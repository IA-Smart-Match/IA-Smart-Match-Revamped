#!/usr/bin/env python3
"""Class-exercise simulation for Dr. Wang: Northline and Harbor, many runs.

Re-run it with:

    python ann_simulation.py                      # 3 rounds x 5,000 runs, seed 1
    python ann_simulation.py --runs 500 --rounds 2 --seed 7
    python ann_simulation.py --self-check         # quick checks, about 1 minute

Needs Python 3.10+ and only the standard library; ``openpyxl`` is used only to
write the spreadsheet (``pip install openpyxl``). The same arguments always give
the same numbers, on any machine (every chance draw is a SHA-256 digest, not
``random``).

What one "run" is
-----------------
One run is one team in class with one chance draw (the team "seed"). Round r
uses seeds ``seed + (r-1)*runs`` up to ``seed + r*runs - 1``, so rounds never
share a seed. In every run, for both list sizes (60 and 30 emails) and four
settings (major first 1/0/0/0, interest first 0/1/0/0, balanced 3/3/2/2, best
possible list), it plays:

1. Northline (event E11) with the setting's list;
2. Harbor (E12) before asking for more information;
3. asking for more (promise / small reward / required), then Harbor again with
   the setting re-run on the updated profiles, with the same chance draws.

"Sign-ups added" is the list's extra sign-ups over the mass email alone in the
same run (Ann's "your list added"). "Attendees added" is the same for show-ups.

Assumptions, each with its source
---------------------------------
Sources: E10 = Ann's email of 2026-10-10; E15 = her two replies of 2026-10-15
(docs/decisions/class-exercise-ann-reply-2026-10-15.md); F16 = her Oct 16
feedback (GitHub issue #413); O2 = her revisions of 2026-10-02; S6 = her
2026-10-06 "Step B" reply to Chau as adopted by owner ruling 2026-10-12;
APP = the app's code (``python/smartmatch_domain/.../exercise/``).

A1  Everyone (all 300) gets the mass email; 20 in 100 notice it. E10 Q2.
A2  Only the listed students also get the CBACH email; 60 in 100 notice it.
    E10 Q2. A listed student notices if either is noticed (about 68 in 100).
A3  Nobody who notices nothing signs up. E10 Q2 / APP simulation.py.
A4  Sign-up chance for a student who noticed: start at 5 in 100; +35 if a TRUE
    interest is an event topic; +15 if the TRUE career goal's topic is an event
    topic; +5 for 1-2 past events, +10 for 3 or more (any past event); +5 if
    the student's major is a target major. Capped to 0..100. S6, confirmed for
    +35 and +5/+10 by E15 (second reply). Highest possible: 70.
A5  Each sign-up attends with chance 65 in 100. E10 Q2/Q4.
A6  The 8 students already signed up attend at 65 in 100 too, from a fixed
    draw. E15 Q2. Seats: 60. Seats empty = 60 - (the 8 who came) - attendees.
A7  No extra luck: the fixed draw alone decides who notices, signs up and
    attends; the same run repeated gives the same answer. E10 Q4.
A8  The list only adds: each student's mass-email outcome is the same with or
    without the list (same draw). E10 Q2 / APP.
A9  Lists are built ONLY from what the app knows: major, the profile card
    (stated interests, career goal) and past events. Hidden "true" columns are
    never used to build a list (O2), except for "best possible", which is an
    oracle list sorted by each student's true sign-up chance (F16 "best
    possible list").
A10 List score = weighted average of four factors, weights normalised to sum 1:
    same major (1/0), stated interest overlap (1 if any stated interest is an
    event topic, O2 4a "count in full"), career goal fit (1/0; Undecided and
    Graduate school point at no topic and earn 0 on every event, O2 4b and the
    2026-10-07 answer), past-event topic overlap (0 / 0.5 / 1 for 0 / 1 / 2+
    related past events, E10 part 2). No card = interest and goal unknown (add
    nothing); no past events = past factor unknown. Default weights 3/3/2/2
    (E10 part 2).
A11 Ties: more information first (card > past events > major only), then
    seniors first, then Ann's fixed tie-break order from the file. O2.
A12 Asking for more goes to every student without a card (230 of 300). F16.
    About 30 / 55 / 80 in 100 of them complete a card for promise / small
    reward / required (half-up rounding: 69 / 127 / 184). Under required, 15 in
    100 of those 230 (35 students) stop responding and notice nothing, mass
    email included (E15 Q3, F16). APP asking.py.
A13 A completed card copies the student's true interests and true career goal.
    Northline attendees gain Northline's topics on their profile for matching.
    These updates never count as past events for the sign-up chance. E15.
A14 "Best possible" after asking keeps the same oracle list, minus students who
    stopped responding.
A15 Emails per round: 60 (F16); 30 is also run, to compare with Ann's 30-email
    table and the earlier reports.
A16 Data: Ann's class file as the app loaded it (300 students, 12 events),
    embedded below without names. Its checksum is in ``DATA["checksum"]``.

This file is a copy of the app's rule, not the app itself. The repository test
``tests/unit/test_ann_simulation_parity.py`` proves the copy gives exactly the
app's lists and outcomes on the same inputs.
"""

# ruff: noqa: E501, B905, RUF005, SIM401, RUF059
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import statistics
import time
from dataclasses import dataclass
from decimal import ROUND_HALF_UP, Decimal
from functools import cache
from pathlib import Path

# ---------------------------------------------------------------------------
# Numbers (A1-A6, A12). Same names as the app's EXERCISE_SIMULATION_COEFFICIENTS.
# ---------------------------------------------------------------------------
BASE_SIGNUP = 0.05
TRUE_FIT_LIFT = 0.50
INTEREST_SHARE_OF_FIT = 0.7  # 0.35 interests, 0.15 career goal
FEW_EVENTS_LIFT, FEW_EVENTS_FROM = 0.05, 1
MANY_EVENTS_LIFT, MANY_EVENTS_FROM = 0.10, 3
SAME_MAJOR_LIFT = 0.05
MASS_NOTICE = 0.20
CBACH_NOTICE = 0.60
ATTEND = 0.65
SEATS, EXISTING = 60, 8
CARD_SHARE = {"promise": 0.30, "reward": 0.55, "required": 0.80}
NON_RESPONDING_SHARE = 0.15
ONE_RELATED_PAST_EVENT = 0.5
NORTHLINE, HARBOR = "E11", "E12"
LIMITS = (60, 30)
FACTORS = ("same_major", "stated_interest_overlap", "career_goal_fit", "past_event_topic_overlap")
SETTINGS = {
    "major": (1.0, 0.0, 0.0, 0.0),
    "interest": (0.0, 1.0, 0.0, 0.0),
    "balanced": (3.0, 3.0, 2.0, 2.0),
    "best": None,
}
SETTING_LABEL = {
    "major": "Major first",
    "interest": "Interest first",
    "balanced": "Balanced 3/3/2/2",
    "best": "Best possible list",
}
CHOICES = ("promise", "reward", "required")
PHASES = ("nl", "before") + tuple(f"after_{c}" for c in CHOICES)
PHASE_LABEL = {
    "nl": "Northline",
    "before": "Harbor, before asking",
    "after_promise": "Harbor, after: promise",
    "after_reward": "Harbor, after: small reward",
    "after_required": "Harbor, after: required",
}
YEAR_RANK = {"Freshman": 1, "Sophomore": 2, "Junior": 3, "Senior": 4}
CAREER_GOAL_TOPICS = {
    "Accounting or audit role (CPA path)": "Accounting / professional services",
    "Consulting role": "Consulting",
    "Media or entertainment business role": "Entertainment / sports / media",
    "Startup or small business role": "Entrepreneurship / startups",
    "Finance or real estate role": "Finance / banking / insurance / real estate",
    "Government or nonprofit role": "Government / nonprofit / public policy",
    "Healthcare administration role": "Healthcare administration",
    "Hospitality or events role": "Hospitality / tourism / events",
    "Marketing or brand role": "Marketing / advertising / public relations",
    "Retail or consumer goods role": "Retail / consumer goods",
    "Sales or business development role": "Sales / business development",
    "Supply chain or operations role": "Supply chain / logistics / operation",
    "Data, analytics or IT role": "Technology / information systems",
    "Start my own business": "Entrepreneurship / startups",
    "Undecided": None,
    "Graduate school": None,
}
# Per-run values recorded for every (emails, phase, setting).
VALUES = (
    "mass_signups",
    "signups",
    "signups_added",
    "mass_attended",
    "attended",
    "attendees_added",
    "existing_attended",
    "seats_empty",
)

# DATA: replaced by the embedded class file (A16). Fields per profile:
# [no, major, class_year, past_event_keys, stated_interests|null, career_goal|null,
#  tiebreak_order, true_interests, true_career_goal|null]; per event:
# [key, name, topics, target_majors].
DATA: dict = json.loads(
    r"""{"checksum":"d6a96df2ef7d064fb112a49ecb291d0250eddf8dbf18b7383a5eaaf28f93eb31","events":[["E01","Resume Lab: Getting Past the Screen",["Sales / business development","Accounting / professional services"],["Accounting","Computer Information Systems","Finance, Real Estate & Law","International Business & Marketing","Management & Human Resources","Technology & Operations Management"]],["E02","Fall Business Career Fair",["Finance / banking / insurance / real estate","Accounting / professional services","Sales / business development"],["Accounting","Computer Information Systems","Finance, Real Estate & Law","International Business & Marketing","Management & Human Resources","Technology & Operations Management"]],["E03","Inside a Big Four Audit Team",["Accounting / professional services","Consulting"],["Accounting"]],["E04","Brand Building at a Streaming Studio",["Marketing / advertising / public relations","Entertainment / sports / media"],["International Business & Marketing"]],["E05","Moving Goods Across the Pacific",["Supply chain / logistics / operation","Retail / consumer goods"],["International Business & Marketing"]],["E06","Pitch Night: Student Startups",["Entrepreneurship / startups","Sales / business development"],["Accounting","Computer Information Systems","Finance, Real Estate & Law","International Business & Marketing","Management & Human Resources","Technology & Operations Management"]],["E07","Careers in Commercial Real Estate",["Finance / banking / insurance / real estate"],["Finance, Real Estate & Law"]],["E08","Cybersecurity and Cloud Jobs Q&A",["Technology / information systems","Consulting"],["Computer Information Systems"]],["E09","Public Service Careers Panel",["Government / nonprofit / public policy","Healthcare administration"],["Accounting","Computer Information Systems","Finance, Real Estate & Law","International Business & Marketing","Management & Human Resources","Technology & Operations Management"]],["E10","Hotel and Event Management Mixer",["Hospitality / tourism / events","Sales / business development"],["Management & Human Resources"]],["E11","Northline Analytics: Behind the Business",["Technology / information systems"],["Computer Information Systems"]],["E12","Harbor Consumer Brands: Behind the Business",["Retail / consumer goods","Supply chain / logistics / operation"],["International Business & Marketing"]]],"profiles":[[1,"Management & Human Resources","Freshman",[],null,null,226,["Hospitality / tourism / events"],"Hospitality or events role"],[2,"Computer Information Systems","Freshman",[],null,null,51,["Finance / banking / insurance / real estate","Technology / information systems"],"Finance or real estate role"],[3,"Accounting","Senior",[],null,null,122,["Accounting / professional services","Retail / consumer goods","Supply chain / logistics / operation"],"Undecided"],[4,"Accounting","Senior",["E07"],["Technology / information systems","Finance / banking / insurance / real estate","Retail / consumer goods"],"Data, analytics or IT role",186,["Technology / information systems","Finance / banking / insurance / real estate","Retail / consumer goods"],"Data, analytics or IT role"],[5,"Finance, Real Estate & Law","Junior",["E10"],null,null,244,["Entertainment / sports / media","Marketing / advertising / public relations"],"Media or entertainment business role"],[6,"Finance, Real Estate & Law","Senior",["E02"],["Sales / business development","Marketing / advertising / public relations","Government / nonprofit / public policy","Accounting / professional services"],"Sales or business development role",120,["Sales / business development","Marketing / advertising / public relations","Government / nonprofit / public policy","Accounting / professional services"],"Sales or business development role"],[7,"Management & Human Resources","Senior",[],null,null,194,["Technology / information systems","Finance / banking / insurance / real estate"],"Data, analytics or IT role"],[8,"Accounting","Senior",[],null,null,2,["Technology / information systems","Finance / banking / insurance / real estate","Accounting / professional services"],"Data, analytics or IT role"],[9,"Management & Human Resources","Junior",[],null,null,111,["Marketing / advertising / public relations","Technology / information systems"],"Marketing or brand role"],[10,"International Business & Marketing","Senior",[],null,null,3,["Government / nonprofit / public policy","Entrepreneurship / startups","Supply chain / logistics / operation"],"Government or nonprofit role"],[11,"Accounting","Freshman",[],null,null,77,["Supply chain / logistics / operation","Accounting / professional services","Technology / information systems"],"Graduate school"],[12,"International Business & Marketing","Senior",[],null,null,227,["Entertainment / sports / media"],"Media or entertainment business role"],[13,"Management & Human Resources","Junior",[],["Hospitality / tourism / events","Government / nonprofit / public policy","Technology / information systems","Entertainment / sports / media"],"Hospitality or events role",230,["Hospitality / tourism / events","Government / nonprofit / public policy","Technology / information systems","Entertainment / sports / media"],"Hospitality or events role"],[14,"Management & Human Resources","Sophomore",[],null,null,140,["Finance / banking / insurance / real estate","Healthcare administration","Accounting / professional services","Marketing / advertising / public relations"],"Finance or real estate role"],[15,"Computer Information Systems","Senior",["E02"],null,null,55,["Hospitality / tourism / events","Sales / business development"],"Hospitality or events role"],[16,"Management & Human Resources","Freshman",["E02"],["Sales / business development","Finance / banking / insurance / real estate","Supply chain / logistics / operation","Consulting"],"Undecided",133,["Sales / business development","Finance / banking / insurance / real estate","Supply chain / logistics / operation","Consulting"],"Undecided"],[17,"Accounting","Junior",["E07"],["Sales / business development","Government / nonprofit / public policy","Finance / banking / insurance / real estate","Accounting / professional services"],"Sales or business development role",207,["Sales / business development","Government / nonprofit / public policy","Finance / banking / insurance / real estate","Accounting / professional services"],"Sales or business development role"],[18,"International Business & Marketing","Senior",["E04"],null,null,37,["Entertainment / sports / media","Marketing / advertising / public relations","Finance / banking / insurance / real estate"],"Undecided"],[19,"International Business & Marketing","Sophomore",["E05","E06","E08"],null,null,251,["Entertainment / sports / media","Finance / banking / insurance / real estate","Sales / business development","Marketing / advertising / public relations"],"Media or entertainment business role"],[20,"Finance, Real Estate & Law","Senior",[],null,null,8,["Hospitality / tourism / events","Government / nonprofit / public policy","Entertainment / sports / media"],"Hospitality or events role"],[21,"Management & Human Resources","Senior",["E04","E06","E08","E10"],null,null,178,["Technology / information systems","Sales / business development","Hospitality / tourism / events","Marketing / advertising / public relations"],"Data, analytics or IT role"],[22,"Management & Human Resources","Senior",[],null,null,104,["Marketing / advertising / public relations","Technology / information systems"],"Marketing or brand role"],[23,"Accounting","Senior",[],null,null,274,["Finance / banking / insurance / real estate","Government / nonprofit / public policy","Entrepreneurship / startups","Marketing / advertising / public relations"],"Finance or real estate role"],[24,"Computer Information Systems","Senior",[],null,null,283,["Technology / information systems","Entrepreneurship / startups"],"Data, analytics or IT role"],[25,"Accounting","Freshman",["E05"],null,null,74,["Supply chain / logistics / operation","Accounting / professional services","Finance / banking / insurance / real estate"],"Supply chain or operations role"],[26,"Management & Human Resources","Junior",["E04"],null,null,124,["Marketing / advertising / public relations","Entertainment / sports / media","Accounting / professional services","Supply chain / logistics / operation"],"Marketing or brand role"],[27,"Finance, Real Estate & Law","Junior",[],null,null,89,["Hospitality / tourism / events","Finance / banking / insurance / real estate","Supply chain / logistics / operation","Entrepreneurship / startups"],"Hospitality or events role"],[28,"Computer Information Systems","Senior",[],["Finance / banking / insurance / real estate","Technology / information systems","Government / nonprofit / public policy"],"Finance or real estate role",149,["Finance / banking / insurance / real estate","Technology / information systems","Government / nonprofit / public policy"],"Finance or real estate role"],[29,"Management & Human Resources","Junior",[],null,null,237,["Retail / consumer goods","Government / nonprofit / public policy","Sales / business development"],"Retail or consumer goods role"],[30,"Management & Human Resources","Sophomore",["E04","E09"],["Sales / business development","Retail / consumer goods","Consulting","Supply chain / logistics / operation"],"Undecided",241,["Sales / business development","Retail / consumer goods","Consulting","Supply chain / logistics / operation"],"Undecided"],[31,"International Business & Marketing","Junior",["E05"],["Retail / consumer goods","Marketing / advertising / public relations","Entertainment / sports / media"],"Retail or consumer goods role",192,["Retail / consumer goods","Marketing / advertising / public relations","Entertainment / sports / media"],"Retail or consumer goods role"],[32,"Management & Human Resources","Sophomore",["E01"],null,null,91,["Sales / business development","Hospitality / tourism / events","Marketing / advertising / public relations"],"Sales or business development role"],[33,"International Business & Marketing","Senior",["E04","E10"],null,null,31,["Entrepreneurship / startups","Entertainment / sports / media","Healthcare administration","Marketing / advertising / public relations"],"Startup or small business role"],[34,"Accounting","Senior",[],null,null,182,["Accounting / professional services","Technology / information systems","Government / nonprofit / public policy"],"Accounting or audit role (CPA path)"],[35,"International Business & Marketing","Junior",["E10"],null,null,234,["Sales / business development","Hospitality / tourism / events","Entertainment / sports / media","Finance / banking / insurance / real estate"],"Sales or business development role"],[36,"Finance, Real Estate & Law","Senior",["E01"],null,null,209,["Finance / banking / insurance / real estate","Consulting","Technology / information systems","Hospitality / tourism / events"],"Undecided"],[37,"Management & Human Resources","Junior",[],null,null,141,["Entrepreneurship / startups"],"Startup or small business role"],[38,"Management & Human Resources","Sophomore",["E04"],["Entertainment / sports / media","Accounting / professional services","Entrepreneurship / startups","Consulting"],"Media or entertainment business role",146,["Entertainment / sports / media","Accounting / professional services","Entrepreneurship / startups","Consulting"],"Media or entertainment business role"],[39,"Accounting","Senior",[],null,null,269,["Finance / banking / insurance / real estate","Marketing / advertising / public relations","Accounting / professional services"],"Finance or real estate role"],[40,"International Business & Marketing","Senior",[],["Marketing / advertising / public relations","Sales / business development","Supply chain / logistics / operation"],"Graduate school",203,["Marketing / advertising / public relations","Sales / business development","Supply chain / logistics / operation"],"Graduate school"],[41,"Computer Information Systems","Senior",[],null,null,212,["Entertainment / sports / media","Supply chain / logistics / operation","Government / nonprofit / public policy","Technology / information systems"],"Media or entertainment business role"],[42,"Finance, Real Estate & Law","Freshman",[],null,null,118,["Accounting / professional services","Government / nonprofit / public policy"],"Accounting or audit role (CPA path)"],[43,"Management & Human Resources","Junior",[],["Hospitality / tourism / events","Marketing / advertising / public relations","Sales / business development","Supply chain / logistics / operation"],"Hospitality or events role",225,["Hospitality / tourism / events","Marketing / advertising / public relations","Sales / business development","Supply chain / logistics / operation"],"Hospitality or events role"],[44,"International Business & Marketing","Senior",["E04","E09"],["Marketing / advertising / public relations","Government / nonprofit / public policy","Finance / banking / insurance / real estate","Accounting / professional services"],"Graduate school",142,["Marketing / advertising / public relations","Government / nonprofit / public policy","Finance / banking / insurance / real estate","Accounting / professional services"],"Graduate school"],[45,"International Business & Marketing","Senior",[],["Entrepreneurship / startups","Hospitality / tourism / events","Sales / business development"],"Startup or small business role",79,["Entrepreneurship / startups","Hospitality / tourism / events","Sales / business development"],"Startup or small business role"],[46,"Accounting","Junior",["E04","E08"],null,null,159,["Consulting","Entertainment / sports / media","Finance / banking / insurance / real estate","Technology / information systems"],"Undecided"],[47,"Management & Human Resources","Junior",[],null,null,294,["Finance / banking / insurance / real estate","Technology / information systems","Hospitality / tourism / events"],"Finance or real estate role"],[48,"Management & Human Resources","Junior",[],null,null,49,["Entrepreneurship / startups"],"Startup or small business role"],[49,"Accounting","Senior",["E01"],["Accounting / professional services","Finance / banking / insurance / real estate","Sales / business development"],"Accounting or audit role (CPA path)",258,["Accounting / professional services","Finance / banking / insurance / real estate","Sales / business development"],"Accounting or audit role (CPA path)"],[50,"International Business & Marketing","Senior",["E04","E05","E09"],["Retail / consumer goods","Entrepreneurship / startups","Marketing / advertising / public relations","Entertainment / sports / media"],"Retail or consumer goods role",280,["Retail / consumer goods","Entrepreneurship / startups","Marketing / advertising / public relations","Entertainment / sports / media"],"Retail or consumer goods role"],[51,"Finance, Real Estate & Law","Junior",["E08"],["Entertainment / sports / media","Finance / banking / insurance / real estate","Government / nonprofit / public policy","Sales / business development"],"Media or entertainment business role",271,["Entertainment / sports / media","Finance / banking / insurance / real estate","Government / nonprofit / public policy","Sales / business development"],"Media or entertainment business role"],[52,"International Business & Marketing","Junior",["E09"],null,null,275,["Sales / business development","Marketing / advertising / public relations","Accounting / professional services","Finance / banking / insurance / real estate"],"Sales or business development role"],[53,"International Business & Marketing","Freshman",[],null,null,205,["Accounting / professional services","Sales / business development","Marketing / advertising / public relations","Entertainment / sports / media"],"Accounting or audit role (CPA path)"],[54,"Accounting","Junior",["E01","E07","E08"],null,null,139,["Entertainment / sports / media","Finance / banking / insurance / real estate","Government / nonprofit / public policy","Accounting / professional services"],"Graduate school"],[55,"International Business & Marketing","Junior",[],null,null,247,["Supply chain / logistics / operation"],"Supply chain or operations role"],[56,"Management & Human Resources","Senior",[],null,null,293,["Marketing / advertising / public relations"],"Undecided"],[57,"Management & Human Resources","Senior",[],null,null,25,["Entertainment / sports / media","Accounting / professional services","Supply chain / logistics / operation"],"Undecided"],[58,"Accounting","Junior",[],["Government / nonprofit / public policy","Accounting / professional services","Retail / consumer goods"],"Government or nonprofit role",145,["Government / nonprofit / public policy","Accounting / professional services","Retail / consumer goods"],"Government or nonprofit role"],[59,"International Business & Marketing","Junior",[],null,null,295,["Marketing / advertising / public relations","Entertainment / sports / media"],"Marketing or brand role"],[60,"Management & Human Resources","Senior",[],null,null,262,["Sales / business development","Supply chain / logistics / operation"],"Sales or business development role"],[61,"Accounting","Junior",["E01","E10"],["Government / nonprofit / public policy","Entrepreneurship / startups","Sales / business development","Finance / banking / insurance / real estate"],"Government or nonprofit role",101,["Government / nonprofit / public policy","Entrepreneurship / startups","Sales / business development","Finance / banking / insurance / real estate"],"Government or nonprofit role"],[62,"Finance, Real Estate & Law","Senior",["E01"],null,null,288,["Retail / consumer goods","Technology / information systems","Supply chain / logistics / operation"],"Undecided"],[63,"Accounting","Sophomore",["E02","E06"],null,null,193,["Finance / banking / insurance / real estate","Sales / business development","Accounting / professional services","Technology / information systems"],"Finance or real estate role"],[64,"International Business & Marketing","Freshman",[],null,null,152,["Marketing / advertising / public relations"],"Start my own business"],[65,"Finance, Real Estate & Law","Freshman",[],null,null,28,["Retail / consumer goods","Marketing / advertising / public relations","Healthcare administration","Finance / banking / insurance / real estate"],"Retail or consumer goods role"],[66,"Management & Human Resources","Junior",[],null,null,279,["Technology / information systems","Supply chain / logistics / operation","Entrepreneurship / startups"],"Undecided"],[67,"Management & Human Resources","Senior",["E08"],null,null,256,["Hospitality / tourism / events","Entertainment / sports / media","Marketing / advertising / public relations"],"Hospitality or events role"],[68,"International Business & Marketing","Junior",["E03"],null,null,197,["Marketing / advertising / public relations","Entertainment / sports / media","Government / nonprofit / public policy","Technology / information systems"],"Marketing or brand role"],[69,"Management & Human Resources","Junior",[],null,null,99,["Government / nonprofit / public policy","Marketing / advertising / public relations"],"Government or nonprofit role"],[70,"International Business & Marketing","Senior",[],null,null,65,["Sales / business development"],"Sales or business development role"],[71,"Management & Human Resources","Senior",[],null,null,17,["Entertainment / sports / media","Sales / business development","Healthcare administration"],"Media or entertainment business role"],[72,"Accounting","Senior",[],["Entertainment / sports / media","Accounting / professional services"],"Media or entertainment business role",157,["Entertainment / sports / media","Accounting / professional services"],"Media or entertainment business role"],[73,"Finance, Real Estate & Law","Senior",[],null,null,198,["Sales / business development","Marketing / advertising / public relations"],"Sales or business development role"],[74,"Finance, Real Estate & Law","Senior",["E04","E09","E10"],null,null,199,["Accounting / professional services","Finance / banking / insurance / real estate","Marketing / advertising / public relations"],"Undecided"],[75,"Technology & Operations Management","Senior",[],["Entertainment / sports / media","Marketing / advertising / public relations"],"Media or entertainment business role",113,["Entertainment / sports / media","Marketing / advertising / public relations"],"Media or entertainment business role"],[76,"International Business & Marketing","Freshman",["E05"],null,null,4,["Finance / banking / insurance / real estate","Accounting / professional services","Entertainment / sports / media","Supply chain / logistics / operation"],"Finance or real estate role"],[77,"Finance, Real Estate & Law","Senior",[],null,null,228,["Sales / business development","Technology / information systems"],"Sales or business development role"],[78,"Management & Human Resources","Senior",[],null,null,268,["Entertainment / sports / media","Entrepreneurship / startups","Finance / banking / insurance / real estate"],"Media or entertainment business role"],[79,"Computer Information Systems","Senior",[],null,null,121,["Finance / banking / insurance / real estate","Marketing / advertising / public relations","Technology / information systems","Government / nonprofit / public policy"],"Finance or real estate role"],[80,"International Business & Marketing","Junior",[],null,null,107,["Marketing / advertising / public relations"],"Marketing or brand role"],[81,"International Business & Marketing","Junior",[],null,null,200,["Entrepreneurship / startups","Marketing / advertising / public relations","Entertainment / sports / media"],"Startup or small business role"],[82,"Finance, Real Estate & Law","Junior",[],null,null,158,["Technology / information systems"],"Data, analytics or IT role"],[83,"Accounting","Senior",[],null,null,239,["Consulting"],"Undecided"],[84,"Management & Human Resources","Junior",["E03","E06"],null,null,170,["Finance / banking / insurance / real estate","Marketing / advertising / public relations","Supply chain / logistics / operation","Entertainment / sports / media"],"Finance or real estate role"],[85,"Management & Human Resources","Sophomore",["E02"],null,null,82,["Accounting / professional services","Entertainment / sports / media","Marketing / advertising / public relations"],"Graduate school"],[86,"Computer Information Systems","Sophomore",[],null,null,78,["Technology / information systems","Finance / banking / insurance / real estate"],"Start my own business"],[87,"Accounting","Sophomore",["E04"],null,null,117,["Sales / business development","Accounting / professional services","Finance / banking / insurance / real estate","Marketing / advertising / public relations"],"Sales or business development role"],[88,"Management & Human Resources","Senior",["E02"],null,null,188,["Technology / information systems","Accounting / professional services","Marketing / advertising / public relations","Sales / business development"],"Data, analytics or IT role"],[89,"Management & Human Resources","Junior",[],null,null,267,["Sales / business development"],"Sales or business development role"],[90,"International Business & Marketing","Junior",[],["Marketing / advertising / public relations","Sales / business development"],"Start my own business",155,["Marketing / advertising / public relations","Sales / business development"],"Start my own business"],[91,"Finance, Real Estate & Law","Senior",[],null,null,291,["Government / nonprofit / public policy"],"Government or nonprofit role"],[92,"Finance, Real Estate & Law","Senior",[],null,null,87,["Sales / business development","Hospitality / tourism / events","Finance / banking / insurance / real estate","Entrepreneurship / startups"],"Sales or business development role"],[93,"International Business & Marketing","Junior",[],null,null,229,["Entertainment / sports / media","Accounting / professional services","Entrepreneurship / startups"],"Undecided"],[94,"Management & Human Resources","Senior",[],null,null,249,["Supply chain / logistics / operation"],"Supply chain or operations role"],[95,"Accounting","Junior",[],["Hospitality / tourism / events","Marketing / advertising / public relations","Accounting / professional services","Government / nonprofit / public policy"],"Hospitality or events role",184,["Hospitality / tourism / events","Marketing / advertising / public relations","Accounting / professional services","Government / nonprofit / public policy"],"Hospitality or events role"],[96,"International Business & Marketing","Senior",[],null,null,253,["Marketing / advertising / public relations","Sales / business development"],"Marketing or brand role"],[97,"Management & Human Resources","Senior",[],null,null,135,["Sales / business development","Technology / information systems","Marketing / advertising / public relations"],"Graduate school"],[98,"Management & Human Resources","Senior",["E03"],["Accounting / professional services","Retail / consumer goods","Supply chain / logistics / operation"],"Start my own business",42,["Accounting / professional services","Retail / consumer goods","Supply chain / logistics / operation"],"Start my own business"],[99,"Computer Information Systems","Senior",[],null,null,21,["Supply chain / logistics / operation","Consulting","Government / nonprofit / public policy","Technology / information systems"],"Supply chain or operations role"],[100,"International Business & Marketing","Junior",[],null,null,6,["Supply chain / logistics / operation","Accounting / professional services","Marketing / advertising / public relations","Finance / banking / insurance / real estate"],"Supply chain or operations role"],[101,"Computer Information Systems","Junior",["E04"],null,null,92,["Supply chain / logistics / operation","Entertainment / sports / media","Marketing / advertising / public relations","Technology / information systems"],"Supply chain or operations role"],[102,"International Business & Marketing","Senior",["E05","E07"],["Entertainment / sports / media","Marketing / advertising / public relations"],"Media or entertainment business role",245,["Entertainment / sports / media","Marketing / advertising / public relations"],"Media or entertainment business role"],[103,"Accounting","Junior",[],null,null,14,["Technology / information systems","Marketing / advertising / public relations"],"Data, analytics or IT role"],[104,"Technology & Operations Management","Junior",["E03"],null,null,41,["Marketing / advertising / public relations","Finance / banking / insurance / real estate","Accounting / professional services"],"Marketing or brand role"],[105,"Management & Human Resources","Senior",["E01"],null,null,85,["Marketing / advertising / public relations","Healthcare administration","Entertainment / sports / media","Accounting / professional services"],"Graduate school"],[106,"Finance, Real Estate & Law","Junior",["E04"],null,null,278,["Retail / consumer goods","Marketing / advertising / public relations","Entertainment / sports / media"],"Retail or consumer goods role"],[107,"Accounting","Junior",[],null,null,224,["Entrepreneurship / startups","Accounting / professional services","Retail / consumer goods"],"Graduate school"],[108,"Computer Information Systems","Senior",["E04"],["Technology / information systems","Supply chain / logistics / operation","Consulting","Government / nonprofit / public policy"],"Data, analytics or IT role",243,["Technology / information systems","Supply chain / logistics / operation","Consulting","Government / nonprofit / public policy"],"Data, analytics or IT role"],[109,"Management & Human Resources","Senior",[],null,null,96,["Finance / banking / insurance / real estate","Retail / consumer goods","Entrepreneurship / startups"],"Finance or real estate role"],[110,"Management & Human Resources","Senior",["E02"],null,null,73,["Sales / business development","Healthcare administration","Marketing / advertising / public relations","Finance / banking / insurance / real estate"],"Undecided"],[111,"Management & Human Resources","Junior",[],null,null,129,["Entertainment / sports / media","Government / nonprofit / public policy","Entrepreneurship / startups","Accounting / professional services"],"Media or entertainment business role"],[112,"Finance, Real Estate & Law","Senior",[],null,null,27,["Sales / business development","Consulting"],"Sales or business development role"],[113,"International Business & Marketing","Junior",[],null,null,23,["Marketing / advertising / public relations","Hospitality / tourism / events"],"Marketing or brand role"],[114,"Computer Information Systems","Senior",[],null,null,189,["Government / nonprofit / public policy","Entrepreneurship / startups","Finance / banking / insurance / real estate","Consulting"],"Graduate school"],[115,"Accounting","Junior",[],null,null,185,["Accounting / professional services","Entrepreneurship / startups","Government / nonprofit / public policy","Supply chain / logistics / operation"],"Accounting or audit role (CPA path)"],[116,"Finance, Real Estate & Law","Junior",[],null,null,282,["Finance / banking / insurance / real estate","Accounting / professional services","Marketing / advertising / public relations"],"Finance or real estate role"],[117,"Accounting","Junior",[],null,null,43,["Supply chain / logistics / operation","Government / nonprofit / public policy","Marketing / advertising / public relations","Entrepreneurship / startups"],"Supply chain or operations role"],[118,"International Business & Marketing","Junior",[],["Finance / banking / insurance / real estate","Marketing / advertising / public relations"],"Finance or real estate role",300,["Finance / banking / insurance / real estate","Marketing / advertising / public relations"],"Finance or real estate role"],[119,"Accounting","Freshman",[],["Entrepreneurship / startups","Finance / banking / insurance / real estate"],"Startup or small business role",235,["Entrepreneurship / startups","Finance / banking / insurance / real estate"],"Startup or small business role"],[120,"Accounting","Senior",[],null,null,181,["Retail / consumer goods","Entrepreneurship / startups","Technology / information systems"],"Retail or consumer goods role"],[121,"Management & Human Resources","Junior",["E02","E03","E04","E06"],null,null,83,["Supply chain / logistics / operation","Finance / banking / insurance / real estate","Marketing / advertising / public relations","Sales / business development"],"Supply chain or operations role"],[122,"Management & Human Resources","Senior",["E07","E08"],["Entertainment / sports / media","Marketing / advertising / public relations","Retail / consumer goods","Entrepreneurship / startups"],"Media or entertainment business role",246,["Entertainment / sports / media","Marketing / advertising / public relations","Retail / consumer goods","Entrepreneurship / startups"],"Media or entertainment business role"],[123,"International Business & Marketing","Senior",[],null,null,5,["Accounting / professional services","Entrepreneurship / startups","Marketing / advertising / public relations","Hospitality / tourism / events"],"Accounting or audit role (CPA path)"],[124,"Computer Information Systems","Junior",["E01","E08","E10"],null,null,137,["Healthcare administration","Technology / information systems","Sales / business development","Hospitality / tourism / events"],"Healthcare administration role"],[125,"Finance, Real Estate & Law","Senior",["E10"],["Finance / banking / insurance / real estate","Accounting / professional services"],"Finance or real estate role",39,["Finance / banking / insurance / real estate","Accounting / professional services"],"Finance or real estate role"],[126,"International Business & Marketing","Junior",[],null,null,7,["Hospitality / tourism / events","Sales / business development","Marketing / advertising / public relations"],"Hospitality or events role"],[127,"Computer Information Systems","Junior",[],null,null,232,["Entertainment / sports / media","Government / nonprofit / public policy"],"Undecided"],[128,"Accounting","Junior",[],null,null,45,["Sales / business development","Supply chain / logistics / operation","Finance / banking / insurance / real estate","Government / nonprofit / public policy"],"Sales or business development role"],[129,"Management & Human Resources","Senior",[],null,null,255,["Marketing / advertising / public relations","Entrepreneurship / startups"],"Graduate school"],[130,"International Business & Marketing","Junior",[],null,null,10,["Accounting / professional services","Government / nonprofit / public policy","Marketing / advertising / public relations"],"Accounting or audit role (CPA path)"],[131,"International Business & Marketing","Freshman",[],["Technology / information systems","Accounting / professional services","Entrepreneurship / startups"],"Data, analytics or IT role",138,["Technology / information systems","Accounting / professional services","Entrepreneurship / startups"],"Data, analytics or IT role"],[132,"Management & Human Resources","Senior",[],["Marketing / advertising / public relations","Technology / information systems"],"Start my own business",143,["Marketing / advertising / public relations","Technology / information systems"],"Start my own business"],[133,"International Business & Marketing","Senior",[],null,null,69,["Supply chain / logistics / operation","Finance / banking / insurance / real estate","Marketing / advertising / public relations","Government / nonprofit / public policy"],"Supply chain or operations role"],[134,"International Business & Marketing","Senior",["E01","E04"],["Entertainment / sports / media","Marketing / advertising / public relations"],"Media or entertainment business role",115,["Entertainment / sports / media","Marketing / advertising / public relations"],"Media or entertainment business role"],[135,"International Business & Marketing","Senior",[],null,null,56,["Entrepreneurship / startups","Finance / banking / insurance / real estate","Marketing / advertising / public relations"],"Startup or small business role"],[136,"Accounting","Senior",[],null,null,195,["Healthcare administration","Entertainment / sports / media"],"Healthcare administration role"],[137,"International Business & Marketing","Junior",[],null,null,156,["Marketing / advertising / public relations","Finance / banking / insurance / real estate","Entrepreneurship / startups"],"Marketing or brand role"],[138,"Accounting","Senior",[],null,null,174,["Entertainment / sports / media","Entrepreneurship / startups"],"Media or entertainment business role"],[139,"Accounting","Senior",[],null,null,148,["Retail / consumer goods","Accounting / professional services","Entertainment / sports / media","Entrepreneurship / startups"],"Retail or consumer goods role"],[140,"Finance, Real Estate & Law","Junior",["E01"],["Accounting / professional services","Finance / banking / insurance / real estate","Marketing / advertising / public relations"],"Accounting or audit role (CPA path)",128,["Accounting / professional services","Finance / banking / insurance / real estate","Marketing / advertising / public relations"],"Accounting or audit role (CPA path)"],[141,"Accounting","Junior",[],null,null,53,["Accounting / professional services","Retail / consumer goods"],"Accounting or audit role (CPA path)"],[142,"Finance, Real Estate & Law","Senior",[],null,null,265,["Supply chain / logistics / operation","Technology / information systems","Finance / banking / insurance / real estate"],"Supply chain or operations role"],[143,"Management & Human Resources","Junior",["E06"],null,null,150,["Retail / consumer goods","Consulting","Supply chain / logistics / operation","Hospitality / tourism / events"],"Retail or consumer goods role"],[144,"Management & Human Resources","Senior",["E04"],null,null,242,["Accounting / professional services","Sales / business development","Marketing / advertising / public relations","Technology / information systems"],"Accounting or audit role (CPA path)"],[145,"International Business & Marketing","Junior",[],["Entertainment / sports / media","Entrepreneurship / startups","Finance / banking / insurance / real estate","Supply chain / logistics / operation"],"Media or entertainment business role",102,["Entertainment / sports / media","Entrepreneurship / startups","Finance / banking / insurance / real estate","Supply chain / logistics / operation"],"Media or entertainment business role"],[146,"Management & Human Resources","Junior",[],null,null,52,["Healthcare administration"],"Healthcare administration role"],[147,"Accounting","Junior",["E01","E02","E04","E05"],null,null,259,["Marketing / advertising / public relations","Entertainment / sports / media"],"Marketing or brand role"],[148,"Accounting","Sophomore",["E03"],null,null,119,["Finance / banking / insurance / real estate","Entertainment / sports / media","Accounting / professional services","Sales / business development"],"Finance or real estate role"],[149,"Accounting","Senior",["E01","E02","E10"],null,null,76,["Sales / business development","Accounting / professional services","Consulting","Hospitality / tourism / events"],"Sales or business development role"],[150,"Accounting","Junior",["E07"],null,null,123,["Finance / banking / insurance / real estate","Retail / consumer goods","Marketing / advertising / public relations","Accounting / professional services"],"Finance or real estate role"],[151,"Accounting","Junior",[],null,null,67,["Accounting / professional services","Entertainment / sports / media","Technology / information systems"],"Accounting or audit role (CPA path)"],[152,"Finance, Real Estate & Law","Freshman",[],null,null,127,["Finance / banking / insurance / real estate"],"Graduate school"],[153,"Accounting","Senior",["E04"],["Sales / business development","Entertainment / sports / media","Hospitality / tourism / events","Government / nonprofit / public policy"],"Sales or business development role",254,["Sales / business development","Entertainment / sports / media","Hospitality / tourism / events","Government / nonprofit / public policy"],"Sales or business development role"],[154,"Finance, Real Estate & Law","Freshman",["E02","E06"],["Retail / consumer goods","Sales / business development","Accounting / professional services","Marketing / advertising / public relations"],"Retail or consumer goods role",180,["Retail / consumer goods","Sales / business development","Accounting / professional services","Marketing / advertising / public relations"],"Retail or consumer goods role"],[155,"Finance, Real Estate & Law","Junior",[],null,null,32,["Entertainment / sports / media","Sales / business development"],"Media or entertainment business role"],[156,"International Business & Marketing","Senior",[],null,null,109,["Marketing / advertising / public relations","Technology / information systems"],"Start my own business"],[157,"Management & Human Resources","Senior",[],null,null,220,["Supply chain / logistics / operation"],"Supply chain or operations role"],[158,"Management & Human Resources","Freshman",[],null,null,177,["Marketing / advertising / public relations"],"Marketing or brand role"],[159,"International Business & Marketing","Senior",[],null,null,131,["Marketing / advertising / public relations","Accounting / professional services","Technology / information systems","Finance / banking / insurance / real estate"],"Marketing or brand role"],[160,"Management & Human Resources","Junior",[],null,null,9,["Finance / banking / insurance / real estate","Government / nonprofit / public policy"],"Finance or real estate role"],[161,"Finance, Real Estate & Law","Senior",[],null,null,112,["Consulting","Healthcare administration","Finance / banking / insurance / real estate","Hospitality / tourism / events"],"Consulting role"],[162,"International Business & Marketing","Senior",[],null,null,190,["Accounting / professional services","Marketing / advertising / public relations","Entertainment / sports / media","Entrepreneurship / startups"],"Start my own business"],[163,"International Business & Marketing","Senior",[],["Entrepreneurship / startups","Marketing / advertising / public relations","Technology / information systems","Accounting / professional services"],"Startup or small business role",136,["Entrepreneurship / startups","Marketing / advertising / public relations","Technology / information systems","Accounting / professional services"],"Startup or small business role"],[164,"Finance, Real Estate & Law","Senior",[],["Finance / banking / insurance / real estate","Accounting / professional services"],"Finance or real estate role",130,["Finance / banking / insurance / real estate","Accounting / professional services"],"Finance or real estate role"],[165,"Management & Human Resources","Sophomore",[],null,null,233,["Entrepreneurship / startups","Healthcare administration"],"Startup or small business role"],[166,"International Business & Marketing","Sophomore",[],null,null,240,["Finance / banking / insurance / real estate","Supply chain / logistics / operation","Marketing / advertising / public relations","Entrepreneurship / startups"],"Finance or real estate role"],[167,"Finance, Real Estate & Law","Senior",[],null,null,214,["Marketing / advertising / public relations","Sales / business development"],"Graduate school"],[168,"Computer Information Systems","Freshman",[],["Supply chain / logistics / operation","Marketing / advertising / public relations","Accounting / professional services"],"Supply chain or operations role",16,["Supply chain / logistics / operation","Marketing / advertising / public relations","Accounting / professional services"],"Supply chain or operations role"],[169,"International Business & Marketing","Freshman",["E04","E06","E07","E08","E10"],["Entertainment / sports / media","Marketing / advertising / public relations","Technology / information systems"],"Media or entertainment business role",35,["Entertainment / sports / media","Marketing / advertising / public relations","Technology / information systems"],"Media or entertainment business role"],[170,"Management & Human Resources","Sophomore",[],null,null,299,["Sales / business development"],"Sales or business development role"],[171,"International Business & Marketing","Senior",[],["Finance / banking / insurance / real estate","Sales / business development","Entrepreneurship / startups","Retail / consumer goods"],"Start my own business",154,["Finance / banking / insurance / real estate","Sales / business development","Entrepreneurship / startups","Retail / consumer goods"],"Start my own business"],[172,"Technology & Operations Management","Senior",[],null,null,179,["Finance / banking / insurance / real estate","Technology / information systems","Government / nonprofit / public policy"],"Finance or real estate role"],[173,"Finance, Real Estate & Law","Freshman",[],null,null,108,["Supply chain / logistics / operation","Accounting / professional services","Hospitality / tourism / events","Sales / business development"],"Graduate school"],[174,"Management & Human Resources","Sophomore",[],null,null,26,["Marketing / advertising / public relations","Government / nonprofit / public policy","Accounting / professional services"],"Start my own business"],[175,"Accounting","Sophomore",[],null,null,24,["Accounting / professional services","Technology / information systems"],"Undecided"],[176,"Finance, Real Estate & Law","Senior",[],null,null,286,["Supply chain / logistics / operation","Consulting","Entrepreneurship / startups"],"Supply chain or operations role"],[177,"Technology & Operations Management","Senior",[],null,null,296,["Entertainment / sports / media","Marketing / advertising / public relations","Healthcare administration"],"Media or entertainment business role"],[178,"Finance, Real Estate & Law","Senior",["E01"],["Technology / information systems","Finance / banking / insurance / real estate","Sales / business development"],"Data, analytics or IT role",261,["Technology / information systems","Finance / banking / insurance / real estate","Sales / business development"],"Data, analytics or IT role"],[179,"Management & Human Resources","Sophomore",[],null,null,219,["Sales / business development","Retail / consumer goods"],"Sales or business development role"],[180,"Finance, Real Estate & Law","Junior",[],null,null,289,["Finance / banking / insurance / real estate"],"Finance or real estate role"],[181,"International Business & Marketing","Senior",[],null,null,285,["Marketing / advertising / public relations","Accounting / professional services","Technology / information systems","Healthcare administration"],"Marketing or brand role"],[182,"Management & Human Resources","Senior",[],null,null,213,["Accounting / professional services","Retail / consumer goods","Technology / information systems"],"Accounting or audit role (CPA path)"],[183,"International Business & Marketing","Senior",["E08"],null,null,80,["Finance / banking / insurance / real estate","Marketing / advertising / public relations","Technology / information systems","Sales / business development"],"Finance or real estate role"],[184,"Finance, Real Estate & Law","Sophomore",[],null,null,48,["Supply chain / logistics / operation","Sales / business development","Entrepreneurship / startups"],"Supply chain or operations role"],[185,"Technology & Operations Management","Senior",[],null,null,211,["Hospitality / tourism / events","Entrepreneurship / startups","Accounting / professional services"],"Hospitality or events role"],[186,"International Business & Marketing","Junior",["E02","E04","E10"],null,null,201,["Supply chain / logistics / operation","Entrepreneurship / startups","Sales / business development","Marketing / advertising / public relations"],"Supply chain or operations role"],[187,"Technology & Operations Management","Senior",[],null,null,126,["Entrepreneurship / startups","Accounting / professional services","Finance / banking / insurance / real estate"],"Start my own business"],[188,"Finance, Real Estate & Law","Junior",["E04"],["Entrepreneurship / startups","Marketing / advertising / public relations","Accounting / professional services","Sales / business development"],"Startup or small business role",187,["Entrepreneurship / startups","Marketing / advertising / public relations","Accounting / professional services","Sales / business development"],"Startup or small business role"],[189,"Management & Human Resources","Senior",["E02"],null,null,270,["Sales / business development","Entrepreneurship / startups","Marketing / advertising / public relations","Finance / banking / insurance / real estate"],"Sales or business development role"],[190,"International Business & Marketing","Junior",["E08"],null,null,50,["Retail / consumer goods","Marketing / advertising / public relations","Technology / information systems","Consulting"],"Retail or consumer goods role"],[191,"Finance, Real Estate & Law","Senior",[],null,null,19,["Accounting / professional services"],"Graduate school"],[192,"Accounting","Junior",["E01","E07"],["Government / nonprofit / public policy","Finance / banking / insurance / real estate","Retail / consumer goods","Accounting / professional services"],"Government or nonprofit role",22,["Government / nonprofit / public policy","Finance / banking / insurance / real estate","Retail / consumer goods","Accounting / professional services"],"Government or nonprofit role"],[193,"Computer Information Systems","Freshman",[],null,null,59,["Government / nonprofit / public policy","Entrepreneurship / startups","Supply chain / logistics / operation"],"Government or nonprofit role"],[194,"International Business & Marketing","Freshman",[],["Hospitality / tourism / events","Retail / consumer goods","Finance / banking / insurance / real estate","Entertainment / sports / media"],"Hospitality or events role",276,["Hospitality / tourism / events","Retail / consumer goods","Finance / banking / insurance / real estate","Entertainment / sports / media"],"Hospitality or events role"],[195,"Finance, Real Estate & Law","Sophomore",[],null,null,132,["Marketing / advertising / public relations"],"Graduate school"],[196,"Finance, Real Estate & Law","Senior",["E10"],null,null,116,["Sales / business development","Entrepreneurship / startups"],"Sales or business development role"],[197,"Accounting","Junior",["E02","E08","E09"],["Marketing / advertising / public relations","Finance / banking / insurance / real estate","Sales / business development"],"Undecided",147,["Marketing / advertising / public relations","Finance / banking / insurance / real estate","Sales / business development"],"Undecided"],[198,"Technology & Operations Management","Senior",[],null,null,252,["Entrepreneurship / startups"],"Startup or small business role"],[199,"Management & Human Resources","Junior",[],null,null,206,["Government / nonprofit / public policy","Entrepreneurship / startups","Entertainment / sports / media","Supply chain / logistics / operation"],"Government or nonprofit role"],[200,"Management & Human Resources","Freshman",[],null,null,260,["Government / nonprofit / public policy","Healthcare administration"],"Graduate school"],[201,"Management & Human Resources","Freshman",[],["Entertainment / sports / media","Hospitality / tourism / events","Finance / banking / insurance / real estate","Marketing / advertising / public relations"],"Media or entertainment business role",167,["Entertainment / sports / media","Hospitality / tourism / events","Finance / banking / insurance / real estate","Marketing / advertising / public relations"],"Media or entertainment business role"],[202,"Accounting","Freshman",[],["Sales / business development","Accounting / professional services"],"Undecided",144,["Sales / business development","Accounting / professional services"],"Undecided"],[203,"Finance, Real Estate & Law","Junior",[],null,null,216,["Technology / information systems","Entrepreneurship / startups"],"Data, analytics or IT role"],[204,"International Business & Marketing","Junior",[],null,null,298,["Supply chain / logistics / operation","Healthcare administration","Marketing / advertising / public relations"],"Start my own business"],[205,"Management & Human Resources","Freshman",[],null,null,236,["Finance / banking / insurance / real estate"],"Finance or real estate role"],[206,"Management & Human Resources","Senior",[],null,null,60,["Finance / banking / insurance / real estate","Entrepreneurship / startups"],"Finance or real estate role"],[207,"Management & Human Resources","Senior",["E06"],null,null,86,["Marketing / advertising / public relations","Sales / business development","Finance / banking / insurance / real estate","Entertainment / sports / media"],"Undecided"],[208,"Computer Information Systems","Senior",["E05"],null,null,46,["Technology / information systems","Consulting","Retail / consumer goods","Supply chain / logistics / operation"],"Graduate school"],[209,"Finance, Real Estate & Law","Junior",[],null,null,81,["Finance / banking / insurance / real estate","Government / nonprofit / public policy","Consulting","Marketing / advertising / public relations"],"Finance or real estate role"],[210,"Finance, Real Estate & Law","Senior",[],null,null,248,["Technology / information systems","Finance / banking / insurance / real estate","Supply chain / logistics / operation","Marketing / advertising / public relations"],"Graduate school"],[211,"Computer Information Systems","Junior",[],null,null,175,["Finance / banking / insurance / real estate","Entrepreneurship / startups","Technology / information systems"],"Finance or real estate role"],[212,"Accounting","Junior",["E10"],null,null,284,["Marketing / advertising / public relations","Accounting / professional services","Sales / business development","Supply chain / logistics / operation"],"Marketing or brand role"],[213,"Computer Information Systems","Junior",[],null,null,169,["Entrepreneurship / startups","Hospitality / tourism / events","Marketing / advertising / public relations","Government / nonprofit / public policy"],"Startup or small business role"],[214,"Management & Human Resources","Sophomore",[],null,null,172,["Technology / information systems","Supply chain / logistics / operation","Accounting / professional services"],"Data, analytics or IT role"],[215,"Management & Human Resources","Senior",["E02","E07","E10"],["Finance / banking / insurance / real estate","Entertainment / sports / media","Marketing / advertising / public relations","Sales / business development"],"Finance or real estate role",223,["Finance / banking / insurance / real estate","Entertainment / sports / media","Marketing / advertising / public relations","Sales / business development"],"Finance or real estate role"],[216,"International Business & Marketing","Junior",[],null,null,281,["Supply chain / logistics / operation","Marketing / advertising / public relations"],"Supply chain or operations role"],[217,"Accounting","Junior",[],null,null,98,["Marketing / advertising / public relations","Finance / banking / insurance / real estate"],"Marketing or brand role"],[218,"International Business & Marketing","Junior",[],null,null,218,["Consulting","Entertainment / sports / media","Entrepreneurship / startups"],"Start my own business"],[219,"Management & Human Resources","Senior",["E05"],["Sales / business development","Supply chain / logistics / operation","Entrepreneurship / startups","Accounting / professional services"],"Sales or business development role",68,["Sales / business development","Supply chain / logistics / operation","Entrepreneurship / startups","Accounting / professional services"],"Sales or business development role"],[220,"Computer Information Systems","Senior",["E08"],null,null,290,["Sales / business development","Technology / information systems","Consulting"],"Undecided"],[221,"Finance, Real Estate & Law","Junior",[],null,null,257,["Entrepreneurship / startups","Finance / banking / insurance / real estate","Marketing / advertising / public relations","Supply chain / logistics / operation"],"Startup or small business role"],[222,"International Business & Marketing","Junior",[],null,null,287,["Entrepreneurship / startups","Marketing / advertising / public relations"],"Undecided"],[223,"Finance, Real Estate & Law","Junior",["E04"],null,null,57,["Government / nonprofit / public policy","Supply chain / logistics / operation","Sales / business development","Hospitality / tourism / events"],"Government or nonprofit role"],[224,"Computer Information Systems","Junior",[],["Technology / information systems","Sales / business development"],"Data, analytics or IT role",173,["Technology / information systems","Sales / business development"],"Data, analytics or IT role"],[225,"Accounting","Junior",["E02"],["Consulting","Finance / banking / insurance / real estate","Sales / business development","Accounting / professional services"],"Graduate school",71,["Consulting","Finance / banking / insurance / real estate","Sales / business development","Accounting / professional services"],"Graduate school"],[226,"Management & Human Resources","Senior",[],null,null,166,["Finance / banking / insurance / real estate"],"Undecided"],[227,"Management & Human Resources","Senior",[],null,null,70,["Entertainment / sports / media","Healthcare administration","Marketing / advertising / public relations"],"Graduate school"],[228,"Technology & Operations Management","Junior",[],null,null,292,["Government / nonprofit / public policy","Technology / information systems","Hospitality / tourism / events","Entertainment / sports / media"],"Government or nonprofit role"],[229,"Accounting","Freshman",[],null,null,66,["Finance / banking / insurance / real estate","Government / nonprofit / public policy"],"Finance or real estate role"],[230,"International Business & Marketing","Sophomore",[],null,null,168,["Marketing / advertising / public relations","Supply chain / logistics / operation","Finance / banking / insurance / real estate","Entrepreneurship / startups"],"Marketing or brand role"],[231,"Accounting","Freshman",[],null,null,161,["Accounting / professional services","Technology / information systems"],"Accounting or audit role (CPA path)"],[232,"Accounting","Senior",[],null,null,165,["Government / nonprofit / public policy","Hospitality / tourism / events","Marketing / advertising / public relations","Technology / information systems"],"Government or nonprofit role"],[233,"International Business & Marketing","Junior",["E01"],["Marketing / advertising / public relations","Entertainment / sports / media"],"Undecided",47,["Marketing / advertising / public relations","Entertainment / sports / media"],"Undecided"],[234,"International Business & Marketing","Junior",[],["Hospitality / tourism / events","Marketing / advertising / public relations","Entrepreneurship / startups"],"Hospitality or events role",61,["Hospitality / tourism / events","Marketing / advertising / public relations","Entrepreneurship / startups"],"Hospitality or events role"],[235,"International Business & Marketing","Sophomore",["E06","E09"],null,null,34,["Marketing / advertising / public relations","Sales / business development","Hospitality / tourism / events"],"Marketing or brand role"],[236,"Accounting","Sophomore",[],null,null,250,["Government / nonprofit / public policy","Entrepreneurship / startups","Finance / banking / insurance / real estate","Sales / business development"],"Government or nonprofit role"],[237,"International Business & Marketing","Sophomore",["E02"],null,null,171,["Sales / business development","Government / nonprofit / public policy","Marketing / advertising / public relations","Finance / banking / insurance / real estate"],"Start my own business"],[238,"International Business & Marketing","Freshman",[],null,null,13,["Marketing / advertising / public relations"],"Marketing or brand role"],[239,"Accounting","Sophomore",[],["Technology / information systems","Accounting / professional services","Hospitality / tourism / events"],"Graduate school",204,["Technology / information systems","Accounting / professional services","Hospitality / tourism / events"],"Graduate school"],[240,"Technology & Operations Management","Freshman",[],null,null,84,["Healthcare administration","Technology / information systems"],"Healthcare administration role"],[241,"Computer Information Systems","Sophomore",[],null,null,134,["Retail / consumer goods","Finance / banking / insurance / real estate","Technology / information systems","Entrepreneurship / startups"],"Retail or consumer goods role"],[242,"Management & Human Resources","Freshman",[],null,null,63,["Sales / business development","Marketing / advertising / public relations","Healthcare administration"],"Sales or business development role"],[243,"Management & Human Resources","Senior",[],null,null,273,["Entrepreneurship / startups","Retail / consumer goods","Entertainment / sports / media","Accounting / professional services"],"Graduate school"],[244,"International Business & Marketing","Junior",[],null,null,238,["Entertainment / sports / media","Marketing / advertising / public relations","Finance / banking / insurance / real estate"],"Media or entertainment business role"],[245,"Accounting","Senior",["E10"],null,null,20,["Marketing / advertising / public relations","Finance / banking / insurance / real estate","Sales / business development"],"Marketing or brand role"],[246,"International Business & Marketing","Senior",[],null,null,103,["Marketing / advertising / public relations","Consulting","Retail / consumer goods","Hospitality / tourism / events"],"Graduate school"],[247,"Technology & Operations Management","Senior",[],null,null,15,["Healthcare administration","Sales / business development"],"Healthcare administration role"],[248,"International Business & Marketing","Senior",["E10"],null,null,100,["Finance / banking / insurance / real estate","Sales / business development","Hospitality / tourism / events","Entertainment / sports / media"],"Finance or real estate role"],[249,"Finance, Real Estate & Law","Freshman",[],["Technology / information systems","Sales / business development","Government / nonprofit / public policy"],"Graduate school",105,["Technology / information systems","Sales / business development","Government / nonprofit / public policy"],"Graduate school"],[250,"International Business & Marketing","Senior",["E04","E10"],null,null,162,["Retail / consumer goods","Accounting / professional services","Marketing / advertising / public relations","Sales / business development"],"Retail or consumer goods role"],[251,"Management & Human Resources","Senior",[],null,null,36,["Hospitality / tourism / events","Technology / information systems","Finance / banking / insurance / real estate","Entrepreneurship / startups"],"Hospitality or events role"],[252,"Accounting","Sophomore",["E01","E03","E10"],null,null,297,["Accounting / professional services","Supply chain / logistics / operation","Technology / information systems","Finance / banking / insurance / real estate"],"Accounting or audit role (CPA path)"],[253,"Management & Human Resources","Junior",[],["Healthcare administration"],"Healthcare administration role",202,["Healthcare administration"],"Healthcare administration role"],[254,"Management & Human Resources","Senior",[],null,null,217,["Accounting / professional services","Technology / information systems","Entrepreneurship / startups","Retail / consumer goods"],"Accounting or audit role (CPA path)"],[255,"Management & Human Resources","Junior",[],null,null,38,["Accounting / professional services","Government / nonprofit / public policy","Marketing / advertising / public relations"],"Accounting or audit role (CPA path)"],[256,"Finance, Real Estate & Law","Sophomore",[],["Retail / consumer goods","Entrepreneurship / startups"],"Start my own business",30,["Retail / consumer goods","Entrepreneurship / startups"],"Start my own business"],[257,"Accounting","Senior",[],null,null,40,["Accounting / professional services","Finance / banking / insurance / real estate","Retail / consumer goods"],"Accounting or audit role (CPA path)"],[258,"Management & Human Resources","Sophomore",["E02"],null,null,44,["Entertainment / sports / media","Marketing / advertising / public relations","Entrepreneurship / startups"],"Start my own business"],[259,"Computer Information Systems","Junior",["E02","E03","E06","E07","E09"],null,null,160,["Entrepreneurship / startups","Hospitality / tourism / events","Finance / banking / insurance / real estate","Sales / business development"],"Undecided"],[260,"International Business & Marketing","Senior",[],null,null,11,["Finance / banking / insurance / real estate","Marketing / advertising / public relations"],"Finance or real estate role"],[261,"Finance, Real Estate & Law","Junior",[],null,null,114,["Entertainment / sports / media","Finance / banking / insurance / real estate"],"Media or entertainment business role"],[262,"Management & Human Resources","Senior",["E02"],["Finance / banking / insurance / real estate","Entrepreneurship / startups","Accounting / professional services","Technology / information systems"],"Finance or real estate role",110,["Finance / banking / insurance / real estate","Entrepreneurship / startups","Accounting / professional services","Technology / information systems"],"Finance or real estate role"],[263,"Accounting","Senior",[],null,null,54,["Entrepreneurship / startups","Healthcare administration","Marketing / advertising / public relations"],"Startup or small business role"],[264,"International Business & Marketing","Sophomore",["E04","E06","E10"],["Sales / business development","Marketing / advertising / public relations","Entertainment / sports / media"],"Start my own business",221,["Sales / business development","Marketing / advertising / public relations","Entertainment / sports / media"],"Start my own business"],[265,"Management & Human Resources","Sophomore",["E04","E06","E08"],["Entertainment / sports / media","Retail / consumer goods","Technology / information systems","Marketing / advertising / public relations"],"Media or entertainment business role",95,["Entertainment / sports / media","Retail / consumer goods","Technology / information systems","Marketing / advertising / public relations"],"Media or entertainment business role"],[266,"Finance, Real Estate & Law","Sophomore",[],null,null,1,["Accounting / professional services","Consulting","Finance / banking / insurance / real estate","Marketing / advertising / public relations"],"Accounting or audit role (CPA path)"],[267,"Finance, Real Estate & Law","Freshman",["E08","E10"],null,null,266,["Technology / information systems","Consulting","Entertainment / sports / media","Sales / business development"],"Graduate school"],[268,"Finance, Real Estate & Law","Senior",[],["Entrepreneurship / startups","Retail / consumer goods","Marketing / advertising / public relations","Accounting / professional services"],"Startup or small business role",272,["Entrepreneurship / startups","Retail / consumer goods","Marketing / advertising / public relations","Accounting / professional services"],"Startup or small business role"],[269,"Finance, Real Estate & Law","Junior",["E04","E08","E09"],null,null,215,["Entertainment / sports / media","Finance / banking / insurance / real estate","Accounting / professional services","Government / nonprofit / public policy"],"Start my own business"],[270,"Finance, Real Estate & Law","Senior",[],null,null,94,["Retail / consumer goods"],"Undecided"],[271,"Finance, Real Estate & Law","Senior",[],null,null,125,["Finance / banking / insurance / real estate","Consulting"],"Finance or real estate role"],[272,"Accounting","Junior",[],["Government / nonprofit / public policy","Retail / consumer goods","Finance / banking / insurance / real estate","Entrepreneurship / startups"],"Government or nonprofit role",196,["Government / nonprofit / public policy","Retail / consumer goods","Finance / banking / insurance / real estate","Entrepreneurship / startups"],"Government or nonprofit role"],[273,"Finance, Real Estate & Law","Junior",["E02"],null,null,90,["Entrepreneurship / startups","Marketing / advertising / public relations","Sales / business development","Finance / banking / insurance / real estate"],"Startup or small business role"],[274,"Management & Human Resources","Junior",[],null,null,75,["Entrepreneurship / startups","Technology / information systems","Accounting / professional services","Marketing / advertising / public relations"],"Startup or small business role"],[275,"International Business & Marketing","Junior",[],null,null,106,["Finance / banking / insurance / real estate","Marketing / advertising / public relations"],"Undecided"],[276,"Management & Human Resources","Senior",["E02","E07"],["Accounting / professional services","Finance / banking / insurance / real estate"],"Accounting or audit role (CPA path)",58,["Accounting / professional services","Finance / banking / insurance / real estate"],"Accounting or audit role (CPA path)"],[277,"Finance, Real Estate & Law","Junior",[],null,null,29,["Hospitality / tourism / events"],"Hospitality or events role"],[278,"Management & Human Resources","Senior",["E01","E03","E07","E10"],null,null,163,["Sales / business development","Accounting / professional services","Technology / information systems","Marketing / advertising / public relations"],"Undecided"],[279,"International Business & Marketing","Senior",[],["Sales / business development","Marketing / advertising / public relations"],"Start my own business",33,["Sales / business development","Marketing / advertising / public relations"],"Start my own business"],[280,"International Business & Marketing","Junior",[],null,null,191,["Accounting / professional services"],"Graduate school"],[281,"Management & Human Resources","Junior",["E09"],null,null,18,["Entertainment / sports / media","Technology / information systems","Supply chain / logistics / operation","Retail / consumer goods"],"Media or entertainment business role"],[282,"Accounting","Freshman",["E02"],null,null,164,["Healthcare administration","Finance / banking / insurance / real estate","Accounting / professional services","Entertainment / sports / media"],"Healthcare administration role"],[283,"Management & Human Resources","Freshman",["E09"],["Entertainment / sports / media","Finance / banking / insurance / real estate","Government / nonprofit / public policy","Marketing / advertising / public relations"],"Media or entertainment business role",64,["Entertainment / sports / media","Finance / banking / insurance / real estate","Government / nonprofit / public policy","Marketing / advertising / public relations"],"Media or entertainment business role"],[284,"Technology & Operations Management","Senior",["E02"],null,null,62,["Entrepreneurship / startups","Sales / business development"],"Undecided"],[285,"Management & Human Resources","Senior",["E05"],null,null,72,["Accounting / professional services","Entrepreneurship / startups","Finance / banking / insurance / real estate","Sales / business development"],"Accounting or audit role (CPA path)"],[286,"Computer Information Systems","Sophomore",[],["Government / nonprofit / public policy","Technology / information systems","Finance / banking / insurance / real estate"],"Government or nonprofit role",183,["Government / nonprofit / public policy","Technology / information systems","Finance / banking / insurance / real estate"],"Government or nonprofit role"],[287,"Technology & Operations Management","Junior",[],null,null,222,["Finance / banking / insurance / real estate","Entrepreneurship / startups","Supply chain / logistics / operation","Technology / information systems"],"Start my own business"],[288,"Management & Human Resources","Junior",[],null,null,264,["Entrepreneurship / startups","Hospitality / tourism / events","Technology / information systems","Government / nonprofit / public policy"],"Startup or small business role"],[289,"Management & Human Resources","Senior",[],["Finance / banking / insurance / real estate","Government / nonprofit / public policy","Accounting / professional services","Supply chain / logistics / operation"],"Undecided",231,["Finance / banking / insurance / real estate","Government / nonprofit / public policy","Accounting / professional services","Supply chain / logistics / operation"],"Undecided"],[290,"Computer Information Systems","Senior",[],null,null,263,["Healthcare administration","Technology / information systems"],"Healthcare administration role"],[291,"Finance, Real Estate & Law","Sophomore",["E05"],["Retail / consumer goods","Finance / banking / insurance / real estate","Supply chain / logistics / operation","Entertainment / sports / media"],"Start my own business",210,["Retail / consumer goods","Finance / banking / insurance / real estate","Supply chain / logistics / operation","Entertainment / sports / media"],"Start my own business"],[292,"Management & Human Resources","Senior",[],null,null,93,["Marketing / advertising / public relations","Entrepreneurship / startups","Supply chain / logistics / operation"],"Marketing or brand role"],[293,"Management & Human Resources","Junior",[],null,null,277,["Marketing / advertising / public relations","Finance / banking / insurance / real estate","Entrepreneurship / startups"],"Marketing or brand role"],[294,"Finance, Real Estate & Law","Junior",[],null,null,151,["Government / nonprofit / public policy","Finance / banking / insurance / real estate","Technology / information systems","Supply chain / logistics / operation"],"Graduate school"],[295,"International Business & Marketing","Junior",[],null,null,12,["Entrepreneurship / startups","Supply chain / logistics / operation"],"Undecided"],[296,"Technology & Operations Management","Freshman",[],null,null,88,["Technology / information systems","Marketing / advertising / public relations","Healthcare administration","Finance / banking / insurance / real estate"],"Data, analytics or IT role"],[297,"Computer Information Systems","Sophomore",[],null,null,208,["Sales / business development","Healthcare administration","Technology / information systems"],"Sales or business development role"],[298,"Accounting","Sophomore",["E02"],null,null,176,["Sales / business development","Marketing / advertising / public relations","Accounting / professional services","Entertainment / sports / media"],"Undecided"],[299,"Computer Information Systems","Senior",[],null,null,97,["Accounting / professional services","Healthcare administration","Technology / information systems","Supply chain / logistics / operation"],"Accounting or audit role (CPA path)"],[300,"International Business & Marketing","Senior",[],["Entrepreneurship / startups","Government / nonprofit / public policy","Finance / banking / insurance / real estate","Sales / business development"],"Startup or small business role",153,["Entrepreneurship / startups","Government / nonprofit / public policy","Finance / banking / insurance / real estate","Sales / business development"],"Startup or small business role"]]}"""
)


# ---------------------------------------------------------------------------
# The rule (copy of the app; see the parity test)
# ---------------------------------------------------------------------------
def norm(term: str) -> str:
    return term.strip().casefold()


def norms(terms) -> frozenset:
    return frozenset(n for n in (norm(t) for t in terms) if n)


_GOAL_BY_NORM = {norm(k): v for k, v in CAREER_GOAL_TOPICS.items()}


def goal_topic(goal: str | None) -> str | None:
    """The topic a career-goal label points at (None for Undecided / Graduate school)."""
    if goal is None:
        return None
    key = norm(goal)
    return _GOAL_BY_NORM[key] if key in _GOAL_BY_NORM else goal


def _text(value: str | None) -> str | None:
    if value is None:
        return None
    return value.strip() or None


def digest(*fields) -> bytes:
    h = hashlib.sha256()
    for f in fields:
        b = str(f).encode()
        h.update(f"{len(b)}:".encode())
        h.update(b)
    return h.digest()


def uniform(seed: int, event_key: str, profile_no: int, purpose: str) -> float:
    return (int.from_bytes(digest(seed, event_key, profile_no, purpose)[:8], "big") >> 11) / 2.0**53


@dataclass(frozen=True)
class Profile:
    no: int
    major: str
    year: str
    past: tuple
    stated: tuple | None
    goal: str | None
    tiebreak: int
    true_interests: tuple
    true_goal: str | None


@dataclass(frozen=True)
class Event:
    key: str
    name: str
    topics: tuple
    majors: tuple


PROFILES = tuple(
    Profile(
        p[0],
        p[1],
        p[2],
        tuple(p[3]),
        None if p[4] is None else tuple(p[4]),
        p[5],
        p[6],
        tuple(p[7]),
        p[8],
    )
    for p in DATA.get("profiles", ())
)
EVENTS = {e[0]: Event(e[0], e[1], tuple(e[2]), tuple(e[3])) for e in DATA.get("events", ())}
NO_CARD = tuple(p.no for p in PROFILES if p.stated is None)


def p_signup(p: Profile, ev: Event) -> float:
    """A4: the sign-up chance of a student who noticed an email (true data)."""
    topics = norms(ev.topics)
    share = 0.0
    if topics & norms(p.true_interests):
        share += INTEREST_SHARE_OF_FIT
    tg = goal_topic(p.true_goal)
    if tg is not None and norm(tg) in topics:
        share += 1.0 - INTEREST_SHARE_OF_FIT
    lift = TRUE_FIT_LIFT * share
    n = len(p.past)
    if n >= FEW_EVENTS_FROM:
        lift += MANY_EVENTS_LIFT if n >= MANY_EVENTS_FROM else FEW_EVENTS_LIFT
    if norm(p.major or "") in norms(ev.majors):
        lift += SAME_MAJOR_LIFT
    return min(1.0, max(0.0, BASE_SIGNUP + lift))


def factor_values(p: Profile, ev: Event, completed: bool, gained: tuple) -> tuple:
    """A10: the four factor values (None = unknown) and the information rank."""
    topics, majors = norms(ev.topics), norms(ev.majors)
    same = 1.0 if p.major.strip().casefold() in majors else 0.0
    if completed:
        interests = p.true_interests
        goal = _text(p.true_goal if p.true_goal is not None else p.goal)
    else:
        interests, goal = p.stated, _text(p.goal)
    if interests is None:
        stated = fit = None
    else:
        stated = 1.0 if norms(interests) & topics else 0.0
        gt = goal_topic(goal)
        fit = 0.0 if gt is None else (1.0 if norm(gt) in topics else 0.0)
    attended = [EVENTS[k].topics if k in EVENTS else () for k in p.past]
    if gained:
        attended.append(gained)
    if not attended:
        past = None
    else:
        related = sum(1 for t in attended if norms(t) & topics)
        past = 0.0 if related == 0 else ONE_RELATED_PAST_EVENT if related == 1 else 1.0
    info = 2 if interests is not None else 1 if attended else 0
    return (same, stated, fit, past), info


def rank(
    event_key: str,
    setting: str,
    limit: int,
    completed=frozenset(),
    gainers=frozenset(),
    gained_topics: tuple = (),
) -> tuple:
    """A10/A11: the setting's list of `limit` profile numbers."""
    raw = SETTINGS[setting]
    total = sum(raw)
    w = tuple(x / total for x in raw)
    ev = EVENTS[event_key]
    keyed = []
    for p in PROFILES:
        vals, info = factor_values(
            p, ev, p.no in completed, gained_topics if p.no in gainers else ()
        )
        score = round(sum(wi * v for wi, v in zip(w, vals) if v is not None), 6)
        keyed.append(((-score, -info, -YEAR_RANK.get(p.year, 0), p.tiebreak), p.no))
    keyed.sort()
    return tuple(no for _, no in keyed[:limit])


@cache
def chances(event_key: str) -> dict:
    return {p.no: p_signup(p, EVENTS[event_key]) for p in PROFILES}


def best_list(event_key: str, limit: int, silent=frozenset()) -> tuple:
    """A9/A14: oracle list by true sign-up chance."""
    pb = chances(event_key)
    return tuple(sorted((n for n in pb if n not in silent), key=lambda n: (-pb[n], n))[:limit])


@cache
def list_before(event_key: str, setting: str, limit: int) -> tuple:
    return best_list(event_key, limit) if setting == "best" else rank(event_key, setting, limit)


def draws(seed: int, event_key: str) -> dict:
    """A7: each student's four fixed draws for one run and one event."""
    return {
        p.no: tuple(
            uniform(seed, event_key, p.no, k)
            for k in ("mass_notice", "cbach_notice", "signup", "attend")
        )
        for p in PROFILES
    }


def existing_attended(seed: int, event_key: str) -> int:
    return sum(
        uniform(seed, event_key, n, "existing_attend") < ATTEND for n in range(1, EXISTING + 1)
    )


def play(event_key: str, listed: tuple, d: dict, existing: int, silent=frozenset()) -> tuple:
    """A1-A8: one event for one run. Returns VALUES and the set of attendees."""
    pb, on = chances(event_key), set(listed)
    mass_su = mass_at = su = at = 0
    attendees = set()
    for no, (m, c, s, a) in d.items():
        if no in silent:
            continue
        signs = s < pb[no]
        mass = m < MASS_NOTICE
        if mass and signs:
            mass_su += 1
            mass_at += a < ATTEND
        if (mass or (no in on and c < CBACH_NOTICE)) and signs:
            su += 1
            if a < ATTEND:
                at += 1
                attendees.add(no)
    return (
        mass_su,
        su,
        su - mass_su,
        mass_at,
        at,
        at - mass_at,
        existing,
        max(0, SEATS - existing - at),
    ), attendees


def select_share(nos, share: float, seed: int, salt: str, of_size: int | None = None) -> tuple:
    """A12: the app's seeded pick of a rounded share of a pool."""
    distinct = sorted(set(nos))
    size = len(distinct) if of_size is None else of_size
    count = int((Decimal(str(float(share))) * size).quantize(Decimal(1), rounding=ROUND_HALF_UP))
    ordered = sorted(distinct, key=lambda n: (digest(seed, salt, n), n))
    return tuple(sorted(ordered[:count]))


def asking(choice: str, seed: int) -> tuple:
    """A12: (card completers, students who stopped responding)."""
    done = select_share(NO_CARD, CARD_SHARE[choice], seed, "card_completion")
    if choice != "required":
        return done, ()
    rest = [n for n in set(NO_CARD) if n not in set(done)]
    return done, select_share(
        rest, NON_RESPONDING_SHARE, seed, "non_responding", of_size=len(set(NO_CARD))
    )


def one_run(seed: int) -> dict:
    """Every scenario for one team seed: {(emails, phase, setting): VALUES} plus asking counts."""
    out = {}
    dn, dh = draws(seed, NORTHLINE), draws(seed, HARBOR)
    en, eh = existing_attended(seed, NORTHLINE), existing_attended(seed, HARBOR)
    plans = {c: asking(c, seed) for c in CHOICES}
    for c, (done, silent) in plans.items():
        out[("asking", c)] = (len(NO_CARD), len(done), len(silent))
    for limit in LIMITS:
        for s in SETTINGS:
            out[(limit, "nl", s)], att = play(NORTHLINE, list_before(NORTHLINE, s, limit), dn, en)
            out[(limit, "before", s)], _ = play(HARBOR, list_before(HARBOR, s, limit), dh, eh)
            for c, (done, silent) in plans.items():
                silent = frozenset(silent)
                if s == "best":
                    lst = best_list(HARBOR, limit, silent)
                else:
                    lst = rank(
                        HARBOR, s, limit, frozenset(done), frozenset(att), EVENTS[NORTHLINE].topics
                    )
                out[(limit, f"after_{c}", s)], _ = play(HARBOR, lst, dh, eh, silent)
    return out


def _chunk(seeds):
    return [one_run(s) for s in seeds]


def run_round(seeds: list, workers: int) -> list:
    if workers <= 1:
        return _chunk(seeds)
    from multiprocessing import Pool

    parts = [seeds[i :: workers * 4] for i in range(workers * 4)]
    with Pool(workers) as pool:
        res = pool.map(_chunk, parts)
    by_seed = {}
    for part, rs in zip(parts, res):
        by_seed.update(zip(part, rs))
    return [by_seed[s] for s in seeds]


# ---------------------------------------------------------------------------
# Statistics and checks
# ---------------------------------------------------------------------------
def stats(xs: list) -> dict:
    q = statistics.quantiles(xs, n=20, method="inclusive")
    return {
        "mean": statistics.fmean(xs),
        "sd": statistics.stdev(xs) if len(xs) > 1 else 0.0,
        "p5": q[0],
        "p50": statistics.median(xs),
        "p95": q[-1],
        "n": len(xs),
    }


def summarise(runs: list) -> dict:
    keys = [k for k in runs[0] if k[0] != "asking"]
    out = {}
    for k in keys:
        for i, v in enumerate(VALUES):
            out[(*k, v)] = stats([r[k][i] for r in runs])
    for c in CHOICES:
        for i, v in enumerate(("asked", "completed", "stopped_responding")):
            out[("asking", c, v)] = stats([r[("asking", c)][i] for r in runs])
    return out


def expected_added(event_key: str, setting: str, limit: int) -> float:
    """Exact expected sign-ups added for a fixed list: (0.68 - 0.20) x sum of chances."""
    both = 1 - (1 - MASS_NOTICE) * (1 - CBACH_NOTICE)
    pb = chances(event_key)
    return (both - MASS_NOTICE) * sum(pb[n] for n in list_before(event_key, setting, limit))


# Ann's numbers (F16 table, Oct 16). (emails, phase, setting) -> sign-ups added.
ANN_CELLS = {
    (60, "nl", "major"): 7.5,
    (60, "nl", "interest"): 5.3,
    (60, "nl", "balanced"): 9.4,
    (60, "nl", "best"): 13.8,
    (60, "before", "major"): 6.7,
    (60, "after_promise", "major"): 6.9,
    (60, "after_reward", "major"): 7.0,
    (60, "after_required", "major"): 6.8,
    (60, "before", "interest"): 7.0,
    (60, "after_promise", "interest"): 10.8,
    (60, "after_reward", "interest"): 13.5,
    (60, "after_required", "interest"): 13.6,
    (60, "before", "balanced"): 9.2,
    (60, "after_promise", "balanced"): 10.1,
    (60, "after_reward", "balanced"): 11.1,
    (60, "after_required", "balanced"): 12.0,
    (60, "before", "best"): 14.9,
    (30, "before", "major"): 3.7,
    (30, "after_promise", "major"): 3.3,
    (30, "after_reward", "major"): 3.1,
    (30, "after_required", "major"): 3.0,
    (30, "before", "interest"): 5.8,
    (30, "after_promise", "interest"): 6.9,
    (30, "after_reward", "interest"): 6.8,
    (30, "after_required", "interest"): 6.9,
    (30, "before", "balanced"): 4.8,
    (30, "after_promise", "balanced"): 6.3,
    (30, "after_reward", "balanced"): 7.3,
    (30, "after_required", "balanced"): 7.8,
    (30, "before", "best"): 8.4,
}
CELL_TOLERANCE = 0.5  # sign-ups; chosen before running, not tuned


def _m(summary, *k):
    return summary[(*k, "signups_added")]["mean"] if len(k) == 3 else summary[k]["mean"]


def requirement_rows(summary: dict, rounds: list, det_ok: bool) -> list:
    """Every rule from REQUIREMENTS, evaluated against the combined results."""
    all_runs = [r for rd in rounds for r in rd]
    rows = []

    def add(rid, rule, quote, source, target, value, status, note=""):
        rows.append(
            {
                "id": rid,
                "rule": rule,
                "quote": quote,
                "source": source,
                "target": target,
                "value": value,
                "status": status,
                "note": note,
            }
        )

    # --- measured on the runs -------------------------------------------------
    sigs = sum(r[(60, "nl", "balanced")][1] for r in all_runs)
    atts = sum(r[(60, "nl", "balanced")][4] for r in all_runs)
    exp_mass = MASS_NOTICE * sum(chances(NORTHLINE).values())
    got_mass = _m(summary, 60, "nl", "balanced", "mass_signups")
    sd_mass = summary[(60, "nl", "balanced", "mass_signups")]["sd"] / len(all_runs) ** 0.5
    exp_k = expected_added(NORTHLINE, "balanced", 60)
    got_k = _m(summary, 60, "nl", "balanced")
    se_k = summary[(60, "nl", "balanced", "signups_added")]["sd"] / len(all_runs) ** 0.5
    ex = _m(summary, 60, "nl", "balanced", "existing_attended")
    ex_se = summary[(60, "nl", "balanced", "existing_attended")]["sd"] / len(all_runs) ** 0.5
    mass_before = _m(summary, 60, "before", "balanced", "mass_signups")
    mass_req = _m(summary, 60, "after_required", "balanced", "mass_signups")
    asked = {c: summary[("asking", c, "completed")]["mean"] for c in CHOICES}
    silent = summary[("asking", "required", "stopped_responding")]["mean"]
    pmax = max(max(chances(e).values()) for e in (NORTHLINE, HARBOR))
    ok = lambda b: "PASS" if b else "FAIL"  # noqa: E731

    add(
        "R01",
        "Mass email to all 300; 20 in 100 notice it",
        "Everyone (all 300): receives usual mass email or post. 20% notice it.",
        "E10 Q2, 2026-10-10",
        f"mass-alone sign-ups = 0.20 x sum of chances = {exp_mass:.2f}",
        f"{got_mass:.2f} (+/- {sd_mass:.2f})",
        ok(abs(got_mass - exp_mass) < 4 * sd_mass + 0.01),
        "Northline, balanced 60; same check holds for any list.",
    )
    add(
        "R02",
        "CBACH email to the listed students; 60 in 100 notice it (about 68 notice one)",
        "Team's matched 30: also receive a personal email from CBACH. 60% notice it.",
        "E10 Q2, 2026-10-10",
        f"added = 0.48 x sum of listed chances = {exp_k:.2f}",
        f"{got_k:.2f} (+/- {se_k:.2f})",
        ok(abs(got_k - exp_k) < 4 * se_k + 0.01),
        "Northline, balanced 60.",
    )
    add(
        "R03",
        "Each sign-up shows up with chance 65 in 100",
        "Rates apply exactly: 20% mass email, 60% CBACH email, 65% show-up.",
        "E10 Q4, 2026-10-10",
        "attended / signed up = 0.65",
        f"{atts / sigs:.4f}",
        ok(abs(atts / sigs - ATTEND) < 0.005),
    )
    add(
        "R04",
        "The 8 already signed up attend at 65 in 100",
        "the 8 pre-registered students should attend 65% of the time, like everyone else.",
        "E15 Q2, 2026-10-15",
        "8 x 0.65 = 5.20 on average",
        f"{ex:.2f} (+/- {ex_se:.2f})",
        ok(abs(ex - 5.2) < 4 * ex_se + 0.01),
    )
    add(
        "R05",
        "No extra luck; same run repeated gives the same answer",
        "No extra luck. The team's fixed chance draw alone decides who notices, signs up and shows up.",
        "E10 Q4, 2026-10-10",
        "identical results on a repeat",
        "identical" if det_ok else "differs",
        ok(det_ok),
    )
    add(
        "R06",
        "Each team has its own chance draw",
        "keep A (each team own chance draw)",
        "E10 Q1, 2026-10-10",
        "different seeds give different draws",
        "every run uses its own seed",
        "PASS",
    )
    add(
        "R07",
        "Students who stop responding notice nothing, mass email included",
        "They notice nothing, mass email included.",
        "F16 decisions table, 2026-10-16; E15 Q3",
        "mass-email-alone sign-ups at Harbor fall under 'required'",
        f"{mass_before:.2f} before -> {mass_req:.2f} under required",
        ok(mass_req < mass_before),
        "They are skipped before any draw (the app's rule; parity test).",
    )
    add(
        "R08",
        "+35 when a true interest is an event topic",
        'The "+35" in your example is the interest-match bonus, not past events',
        "E15 second reply, 2026-10-15",
        "+35",
        f"+{TRUE_FIT_LIFT * INTEREST_SHARE_OF_FIT * 100:.0f}",
        ok(round(TRUE_FIT_LIFT * INTEREST_SHARE_OF_FIT, 6) == 0.35),
    )
    add(
        "R09",
        "Past events: +5 for 1-2, +10 for 3 or more",
        "past events add +5 (1–2 events) or +10 (3 or more), which is fine.",
        "E15 second reply, 2026-10-15",
        "+5 / +10",
        f"+{FEW_EVENTS_LIFT * 100:.0f} / +{MANY_EVENTS_LIFT * 100:.0f}",
        ok(
            (FEW_EVENTS_LIFT, MANY_EVENTS_LIFT, FEW_EVENTS_FROM, MANY_EVENTS_FROM)
            == (0.05, 0.10, 1, 3)
        ),
    )
    add(
        "R10",
        "Profile updates during the exercise never count as past events",
        "Agree that profile updates during the exercise don't count as past events.",
        "E15 second reply, 2026-10-15",
        "sign-up chance reads the file's past events only",
        "p_signup reads len(file past events) only",
        "PASS",
    )
    add(
        "R11",
        "Start at 5; +15 for a career-goal fit; +5 for same major",
        "(team summary of Ann's Step B) Start at 5 ... +15 if the hidden true career goal fits ... +5 if the "
        "event targets the student's major",
        "S6 2026-10-06 via owner ruling 2026-10-12 (not quoted by Ann since)",
        "5 / +15 / +5",
        f"{BASE_SIGNUP * 100:.0f} / +{TRUE_FIT_LIFT * (1 - INTEREST_SHARE_OF_FIT) * 100:.0f} / "
        f"+{SAME_MAJOR_LIFT * 100:.0f}",
        ok(
            (BASE_SIGNUP, round(TRUE_FIT_LIFT * (1 - INTEREST_SHARE_OF_FIT), 6), SAME_MAJOR_LIFT)
            == (0.05, 0.15, 0.05)
        ),
        "Conflict C16: Ann has not restated these three numbers in her own words.",
    )
    add(
        "R12",
        "Highest sign-up chance is 70",
        "the most anyone can reach is 70 (team summary)",
        "S6 / docs/decisions/class-exercise-current-rule.md",
        "at most 0.70",
        f"highest in the class file: {pmax:.2f}",
        ok(pmax <= 0.70 + 1e-9),
    )
    add(
        "R13",
        "A matching stated interest counts in full",
        "a matching interest should count in full.",
        "O2 item 4a, 2026-10-02",
        "1 for any overlap",
        "1 for any overlap (A10)",
        "PASS",
    )
    add(
        "R14",
        "Past-event topic factor is 0 / 0.5 / 1",
        "Past events 0 / 0.5 / 1: ok",
        "E10 part 2, 2026-10-10",
        "0 / 0.5 / 1",
        f"0 / {ONE_RELATED_PAST_EVENT} / 1",
        ok(ONE_RELATED_PAST_EVENT == 0.5),
    )
    und = factor_values(
        Profile(
            0, "x", "Senior", (), ("Retail / consumer goods",), "Undecided", 0, (), "Undecided"
        ),
        EVENTS[HARBOR],
        False,
        (),
    )[0][2]
    und_p = p_signup(Profile(0, "x", "Senior", (), None, None, 0, (), "Undecided"), EVENTS[HARBOR])
    add(
        "R15",
        "An Undecided career goal earns 0, on every event (Harbor too)",
        "That rule is not in the plan, and Northline is not a broad event. Please remove it.",
        "O2 item 4b, 2026-10-02; answer of 2026-10-07",
        "0 in matching and in the sign-up chance",
        f"matching {und}; sign-up chance {und_p:.2f} (= the 5 start)",
        ok(und == 0.0 and abs(und_p - 0.05) < 1e-9),
    )
    add(
        "R16",
        "Default weights 3/3/2/2",
        "Weights 3,3,2,2, whole steps 0-10, total whole number: ok",
        "E10 part 2, 2026-10-10; ruling 2026-10-06",
        "3/3/2/2",
        "/".join(f"{x:g}" for x in SETTINGS["balanced"]),
        ok(SETTINGS["balanced"] == (3.0, 3.0, 2.0, 2.0)),
    )
    add(
        "R22",
        "Hidden true columns never used to build a list",
        "The app must never show them and never use them to build a list.",
        "O2, 2026-10-02",
        "lists use only major, card, past events",
        "only 'best possible' (an oracle for comparison) reads them; card completers copy them, as the app does",
        "PASS",
    )
    lens = {len(list_before(e, s, 60)) for e in (NORTHLINE, HARBOR) for s in SETTINGS}
    add(
        "R23",
        "60 emails in both rounds",
        "Send 60 tailored, personalized emails in both rounds, not 30.",
        "F16, 2026-10-16",
        "60 per list",
        f"{sorted(lens)}",
        ok(lens == {60}),
    )
    add(
        "R25",
        "The 300-student mass email stays as the comparison",
        "The 300-student mass email stays as the comparison.",
        "F16, 2026-10-16",
        "all 300 get it",
        f"{len(PROFILES)} students",
        ok(len(PROFILES) == 300),
    )
    seats = _m(summary, 60, "nl", "balanced", "seats_empty")
    add(
        "R26",
        "Seats empty out of 60, counting the 8",
        "seats still empty out of 60, counting the 8 already signed up.",
        "checklist row 60 (Ann's test step)",
        "60 - the 8 who came - attendees",
        f"{seats:.1f} empty on average (Northline, balanced 60)",
        "PASS",
    )
    add(
        "R27",
        "Lift line: mass email alone X; with your list Y; your list added Y - X",
        "Mass email alone: 11 sign-ups. With your 30 CBACH emails: 17. Your list added 6.",
        "E10 Q2, 2026-10-10",
        "added = with list - mass alone",
        "computed exactly that way",
        "PASS",
    )
    gap = _m(summary, 60, "nl", "balanced") - _m(summary, 60, "nl", "balanced", "attendees_added")
    add(
        "R28",
        "Attendees added are fewer than sign-ups added; about 3 fewer is fine",
        "If it shows attendance, the ~3 fewer attendees is fine.",
        "E15 Q2, 2026-10-15",
        "about 3 fewer",
        f"{gap:.1f} fewer (Northline, balanced 60)",
        ok(2.0 <= gap <= 4.5),
    )
    add(
        "R32",
        "Asking goes to all students without a card (230)",
        "The card request goes to all students without a card (230 of 300), not only the team's invitees.",
        "F16, 2026-10-16",
        "230",
        f"{len(NO_CARD)}",
        ok(len(NO_CARD) == 230),
    )
    add(
        "R33",
        "About 30 / 55 / 80 in 100 complete a card",
        "about 30%, 55%, and 80% of invited profiles without a card complete one.",
        "checklist row 76 (Ann's test step)",
        "69 / 127 / 184 of 230",
        " / ".join(f"{asked[c]:.0f}" for c in CHOICES),
        ok(tuple(round(asked[c]) for c in CHOICES) == (69, 127, 184)),
    )
    add(
        "R34",
        "Required: about 15 in 100 stop responding",
        "Under 'required,' about 15% of those without a card are marked 'stopped responding,'",
        "checklist row 77 (Ann's test step)",
        "35 of 230",
        f"{silent:.0f}",
        ok(round(silent) == 35),
    )
    add(
        "R38",
        "Before and after asking use the same draws",
        "the gap between the two rows can only come from better information.",
        "F16, 2026-10-16",
        "same seed for both rows",
        "same seed for both rows",
        "PASS",
    )
    add(
        "R47",
        "5,000-run check for both events with 60",
        "a rerun of the 5,000-run check for both events with 60.",
        "F16, 2026-10-16",
        "5,000 runs, both events",
        f"{len(rounds)} rounds x {len(rounds[0]):,} runs",
        ok(all(len(r) >= 5000 for r in rounds)),
    )
    add(
        "R52",
        "'13 already coming' is not all 8 + 5: the 8 attend at 65%",
        "equals all 8 plus the mass email's 5, so please confirm this is live.",
        "F16, 2026-10-16",
        "8 at 65% (about 5.2 of 8 come)",
        f"{ex:.2f} of 8 come",
        ok(abs(ex - 5.2) < 4 * ex_se + 0.01),
        "Live: same rule as R04. Her test run predates the Oct 16 change.",
    )

    # --- Ann's per-cell numbers and targets -----------------------------------
    for k, ann in ANN_CELLS.items():
        L, ph, s = k
        ours = _m(summary, *k)
        rid = f"R42-{s}" if ph == "nl" else (f"R43-{s}-{ph}" if L == 60 else f"R44-{s}-{ph}")
        src = (
            "F16 'Round one with 60', 2026-10-16"
            if ph == "nl"
            else "F16 'Why 60 emails, in numbers' table, 2026-10-16"
        )
        add(
            rid,
            f"{SETTING_LABEL[s]}, {PHASE_LABEL[ph]}, {L} emails: sign-ups added",
            "These are my estimates ... from a simplified version of the rule on the Oct 2 class file"
            if ph != "nl"
            else "The list adds about 7.5 (major first), 5.3 (interest first) and 9.4 (balanced)",
            src,
            f"{ann} (within {CELL_TOLERANCE})",
            f"{ours:.2f} (diff {ours - ann:+.2f})",
            ok(abs(ours - ann) <= CELL_TOLERANCE),
        )
    b = {ph: _m(summary, 60, ph, "balanced") for ph in PHASES}
    add(
        "R45a",
        "Northline target about 9 (balanced, 60)",
        "My rough new targets are about 9 at Northline",
        "F16, 2026-10-16",
        "8.5 to 9.5",
        f"{b['nl']:.2f}",
        ok(8.5 <= b["nl"] <= 9.5),
    )
    add(
        "R45b",
        "Harbor before asking about 9 (balanced, 60)",
        "at Harbor about 9 before asking",
        "F16, 2026-10-16",
        "8.5 to 9.5",
        f"{b['before']:.2f}",
        ok(8.5 <= b["before"] <= 9.5),
    )
    for c in CHOICES:
        v = b[f"after_{c}"]
        add(
            f"R45c-{c}",
            f"Harbor after asking ({c}) 10 to 12 (balanced, 60)",
            "and 10 to 12 after.",
            "F16, 2026-10-16",
            "10 to 12",
            f"{v:.2f}",
            ok(10.0 <= v <= 12.0),
        )
    add(
        "R46",
        "Old 30-email targets (about 6 Northline, 4-5 Harbor) no longer apply",
        "The old targets ('about 6' at Northline, '4 to 5' at Harbor) no longer apply.",
        "F16, 2026-10-16",
        "superseded",
        f"for reference: {_m(summary, 30, 'nl', 'balanced'):.2f} / "
        f"{_m(summary, 30, 'before', 'balanced'):.2f}",
        "N/A",
    )
    # Rounds agree within sampling noise.
    worst = 0.0
    for k in ANN_CELLS:
        ms = [summarise_cell(rd, k) for rd in rounds]
        se = max(s for _, s in ms)
        worst = max(worst, (max(m for m, _ in ms) - min(m for m, _ in ms)) / (se * 2**0.5))
    add(
        "V1",
        "The rounds agree within sampling noise (our own check)",
        "",
        "brief",
        "largest gap < 4 standard errors",
        f"{worst:.2f} standard errors",
        ok(worst < 4.0),
    )

    # --- rules the simulation cannot check (screens, process) -----------------
    for rid, rule, quote, source in UI_RULES:
        add(
            rid,
            rule,
            quote,
            source,
            "-",
            "-",
            "N/A",
            "A screen or process rule; checked on the class checklist, not by the simulation.",
        )
    return rows


def summarise_cell(runs, k):
    xs = [r[k][2] for r in runs]
    return statistics.fmean(xs), statistics.stdev(xs) / len(xs) ** 0.5


UI_RULES = (
    (
        "R17",
        "P004 drops but stays above a plain Accounting major",
        "P004 Brandon Soto: OK (keep, reword checklist",
        "E10 part 2",
    ),
    (
        "R18",
        "Tie order: more information, then seniors, then a fixed order",
        "Works as written: more information first, then seniors first, then a fixed order.",
        "O2",
    ),
    (
        "R19",
        "Years run Freshman to Senior on 'who is on the list'",
        "years should run Freshman, Sophomore, Junior, Senior, not alphabetically.",
        "O2 4d",
    ),
    (
        "R20",
        "One reason line per name",
        "A profile known by major alone says 'same major; nothing else on file.'",
        "O2 4c; F16",
    ),
    (
        "R21",
        "Event description above the sliders",
        "Show a short description at the top of the Northline and Harbor pages",
        "O2 item 5",
    ),
    (
        "R24",
        "60 as an instructor setting (checklist said 5 to 30)",
        "whether 60 is a simple instructor setting",
        "F16 (open question)",
    ),
    (
        "R29",
        "Debrief sentence on the results screen",
        "Each team's numbers include its own luck.",
        "E10 Q1; E15 Q4",
    ),
    ("R30", "One results run per team per event", "A second run must be blocked", "O2; F16"),
    (
        "R31",
        "Lock round two until Northline results",
        "Lock round two, and Harbor results, until round one is done.",
        "F16",
    ),
    (
        "R35",
        "Three asking options, one plain sentence each",
        "Three options each have one plain sentence",
        "checklist row 69",
    ),
    (
        "R36",
        "Way of asking picked after Northline results",
        "Your team picks a way of asking after it has its results for Northline.",
        "E10 addition 2026-10-12",
    ),
    (
        "R37",
        "Major-only setting kept in round two",
        "A major-only setting cannot benefit from new cards. That is a lesson",
        "F16",
    ),
    (
        "R39",
        "Before/after rows do not use up the run",
        "These extra rows do not use up the team's one results run.",
        "F16",
    ),
    ("R40", "Harbor debrief line", "Same event, same students, same luck.", "F16"),
    (
        "R41",
        "'How much we know' table in both rounds",
        "Show it in round one, and again in round two as before and after.",
        "F16",
    ),
    (
        "R48",
        "Chart caption",
        "Light bar: mass email alone. Dark bar: mass email plus your CBACH emails.",
        "F16 results item 5",
    ),
    (
        "R49",
        "Leave 'invited' out of the chart",
        "Leave 'invited' out of the chart, or give it its own chart.",
        "F16 results item 4",
    ),
    (
        "R50",
        "One 'added' measure in the headline",
        "Use one measure in the headline, or label both",
        "F16 results item 1",
    ),
    (
        "R51",
        "Round-one recap matches the round-one screen",
        "Same numbers and wording in both places.",
        "F16 results item 3",
    ),
    (
        "R53",
        "Show 'You will send' with the 60 names",
        "Show 'You will send: [name]' with a short summary of its 60 names",
        "F16 results item 7",
    ),
    (
        "R54",
        "Confirm wording names the setting",
        "Send 'Balanced'? You get one results run for Northline.",
        "F16 results item 8",
    ),
    ("R55", "Round banners", "Round 1 of 2: Northline", "F16"),
    ("R56", "Round-one done message", "Round 1 done. Next: Round 2, Harbor.", "F16"),
    (
        "R57",
        "Nothing from round two before round one results",
        "Show nothing from round two ... until round one results are in.",
        "F16",
    ),
    (
        "R58",
        "Same-topic note",
        "Keep it, but say it only counts toward events on the same topic.",
        "F16 asking item 2",
    ),
    (
        "R59",
        "'Stopped responding' in plain words",
        "Say this in plain words on the screen.",
        "F16 asking item 1",
    ),
    (
        "R60",
        "Clear test teams, close both events",
        "Clear all test teams and close both events before testing week",
        "F16",
    ),
    (
        "R61",
        "Fixed copy for testing and class",
        "Run testing and class from a fixed copy that nobody edits.",
        "F16",
    ),
    ("R62", "Two teams at once, isolated", "Neither should see or change the other's work.", "F16"),
    (
        "R63",
        "Readable on the projector",
        "Readable from the back of the room on the projector.",
        "checklist row 63",
    ),
    (
        "R64",
        "Pages load in under 5 seconds",
        "Each page loads in under 5 seconds in Chrome",
        "checklist row 95",
    ),
    (
        "R65",
        "Wrong passcode message",
        "A wrong passcode shows 'Wrong passcode' and nothing else.",
        "checklist row 11",
    ),
    (
        "R66",
        "Missing-column upload message",
        "You see a plain message naming the missing column",
        "checklist row 13",
    ),
    (
        "R67",
        "List limit can be changed",
        "Change the list limit from 30 to 25. You see 'Limit set to 25.'",
        "checklist row 15",
    ),
    ("R68", "Refresh needs two presses", "Team's own refresh two presses: OK", "E10 part 2"),
    ("R69", "'Refreshed at' time stamp", "must say 'Refreshed at 10:42 AM' too", "E10 part 2"),
    ("R70", "Pinned team status line", "pin it — stays on screen while scrolling", "E10 part 2"),
    ("R71", "New data file gives a new draw", "New file gives new chance draw.", "E10 Q6"),
    ("R72", "License line", "All student profiles are fictional.", "E10 addition 2026-10-12"),
    (
        "R73",
        "Download leaves out the pink columns",
        "The two pink columns from the data file do not appear.",
        "checklist row 51",
    ),
)

CONFLICTS = (
    (
        "C1",
        "Notice, show-up and luck",
        "Sep 25: start 4, +40 'a lot', +/-10 luck, 75 attend. Oct 6 Step A: 75/15 notice, 70 attend, small luck.",
        "Oct 10: 20 / 60 notice, 65 show up, no luck",
        "Used Oct 10",
    ),
    (
        "C2",
        "Sign-up points",
        "Sep 25 translation +40 / +10 / +4. Oct 6 Step B +35 / +15 / +5 or +10 / +5.",
        "Step B; Oct 15 confirms +35 and +5/+10",
        "Used Step B",
    ),
    (
        "C3",
        "Past-event cap",
        "Oct 15 first reply: +5 per event, cap +15. Same-day second reply withdrew it.",
        "Oct 15 second reply and Oct 16: +5 / +10",
        "Used +5 / +10",
    ),
    (
        "C4",
        "'Past events' means two things",
        "List factor counts RELATED past events (0 / 0.5 / 1); sign-up chance counts ANY past event (+5 / +10).",
        "Both approved, for different parts",
        "Kept both",
    ),
    (
        "C5",
        "Undecided career goal",
        "Sep 25: half credit on broad events. Oct 2: remove on Northline. Oct 7: 0 on every event.",
        "Oct 7",
        "Used 0 everywhere",
    ),
    (
        "C6",
        "Default weights",
        "Oct 2: all start at 0.25. Oct 6: 3/3/2/2.",
        "Oct 10 'ok' to 3/3/2/2",
        "Used 3/3/2/2",
    ),
    (
        "C7",
        "List size",
        "Oct 2 and Oct 10: 30. Oct 16: 60. Checklist: 5 to 30.",
        "Oct 16: 60 (instructor-setting range still open)",
        "Ran 60, and 30 for comparison",
    ),
    (
        "C8",
        "Targets",
        "Oct 10: about 6 Northline, 4-5 Harbor (30 emails).",
        "Oct 16: about 9; Harbor about 9 then 10-12",
        "Used Oct 16",
    ),
    (
        "C10",
        "Who is asked for more",
        "Older: only the team's invitees, end of round one (checklist row 76 still says 'invited').",
        "Oct 16: all 230 without a card",
        "Used all 230",
    ),
    (
        "C14",
        "'13 already coming'",
        "Her Oct 16 test run showed all 8 plus 5.",
        "Oct 15: the 8 attend at 65%",
        "Used 65%",
    ),
    (
        "C16",
        "Start 5 / +15 goal / +5 major",
        "Only in the team's summary of her Oct 6 Step B; 'the team's existing sign-up chances' (Oct 10) was read as Step B.",
        "Not restated by Ann",
        "Used Step B; ask Ann to confirm",
    ),
)


# ---------------------------------------------------------------------------
# Output: CSV, JSON, Markdown, spreadsheet
# ---------------------------------------------------------------------------
METRIC_ORDER = [(L, ph, s) for L in LIMITS for ph in PHASES for s in SETTINGS]
SHOWN = (
    "signups_added",
    "attendees_added",
    "mass_signups",
    "signups",
    "attended",
    "existing_attended",
    "seats_empty",
)
VALUE_LABEL = {
    "signups_added": "Sign-ups added by the list",
    "attendees_added": "Attendees added by the list",
    "mass_signups": "Sign-ups, mass email alone",
    "signups": "Sign-ups, mass email plus list",
    "attended": "Attended (not counting the 8)",
    "existing_attended": "Of the 8 already signed up, came",
    "seats_empty": "Seats still empty (of 60)",
    "mass_attended": "Attended, mass email alone",
    "nonresponder_signups": "Sign-ups from students who stopped responding",
}


def write_csv(path: Path, seeds: list, runs: list) -> None:
    with path.open("w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["seed"] + [f"{L}|{ph}|{s}|{v}" for (L, ph, s) in METRIC_ORDER for v in VALUES])
        for seed, r in zip(seeds, runs):
            w.writerow([seed] + [x for k in METRIC_ORDER for x in r[k]])


def write_markdown(path: Path, rows: list, summaries: list, combined: dict, meta: dict) -> None:
    counts = {s: sum(1 for r in rows if r["status"] == s) for s in ("PASS", "FAIL", "N/A")}
    lines = [
        f"# Requirements check: class-exercise simulation ({meta['date']})",
        "",
        f"{counts['PASS']} PASS, {counts['FAIL']} FAIL, {counts['N/A']} not checkable by simulation "
        f"(screens and process). {meta['rounds']} rounds x {meta['runs']:,} runs, seeds {meta['seed_ranges']}.",
        "",
        "Re-run: `python ann_simulation.py --runs 5000 --rounds 3 --seed 1`.",
        "",
        "## Sign-ups added, 60 emails (mean of all rounds; Ann's number in brackets)",
        "",
        "| Setting | Northline | Harbor before | After: promise | After: small reward | After: required |",
        "|---|---|---|---|---|---|",
    ]
    for s in SETTINGS:
        cells = []
        for ph in PHASES:
            m = combined[(60, ph, s, "signups_added")]["mean"]
            ann = ANN_CELLS.get((60, ph, s))
            cells.append(f"{m:.1f}" + (f" [{ann}]" if ann is not None else ""))
        lines.append(f"| {SETTING_LABEL[s]} | " + " | ".join(cells) + " |")
    lines += [
        "",
        "## Per round (balanced 3/3/2/2, 60 emails, sign-ups added: mean / SD)",
        "",
        "| Round | Northline | Harbor before | After: promise | After: small reward | After: required |",
        "|---|---|---|---|---|---|",
    ]
    for i, sm in enumerate(summaries, 1):
        lines.append(
            f"| {i} | "
            + " | ".join(
                f"{sm[(60, ph, 'balanced', 'signups_added')]['mean']:.2f} / {sm[(60, ph, 'balanced', 'signups_added')]['sd']:.2f}"
                for ph in PHASES
            )
            + " |"
        )
    lines += [
        "",
        "## Every rule",
        "",
        "| ID | Rule | Ann's words | Source | Target | Ours | Result | Note |",
        "|---|---|---|---|---|---|---|---|",
    ]
    esc = lambda t: str(t).replace("|", "/")  # noqa: E731
    for r in rows:
        lines.append(
            "| "
            + " | ".join(
                esc(r[k])
                for k in ("id", "rule", "quote", "source", "target", "value", "status", "note")
            )
            + " |"
        )
    lines += [
        "",
        "## Conflicts between sources (newest from Ann wins)",
        "",
        "| ID | Topic | Older | Newest | What we used |",
        "|---|---|---|---|---|",
    ]
    lines += ["| " + " | ".join(esc(x) for x in c) + " |" for c in CONFLICTS]
    path.write_text("\n".join(lines) + "\n")


def write_xlsx(path: Path, rows: list, summaries: list, combined: dict, meta: dict) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    head, bold, base = (
        Font(name="Arial", bold=True, color="FFFFFF"),
        Font(name="Arial", bold=True),
        Font(name="Arial"),
    )
    fill = PatternFill("solid", fgColor="1F4E78")
    status_fill = {
        "PASS": PatternFill("solid", fgColor="C6EFCE"),
        "FAIL": PatternFill("solid", fgColor="FFC7CE"),
        "N/A": PatternFill("solid", fgColor="EDEDED"),
    }

    def table(ws, start_row, header, data, widths=None):
        for j, h in enumerate(header, 1):
            c = ws.cell(row=start_row, column=j, value=h)
            c.font, c.fill, c.alignment = head, fill, Alignment(wrap_text=True, vertical="top")
        for i, row in enumerate(data, start_row + 1):
            for j, v in enumerate(row, 1):
                c = ws.cell(row=i, column=j, value=v)
                c.font, c.alignment = base, Alignment(wrap_text=True, vertical="top")
                if isinstance(v, float):
                    c.number_format = "0.00"
                if v in status_fill:
                    c.fill = status_fill[v]
        for j, wdt in enumerate(widths or [], 1):
            ws.column_dimensions[ws.cell(row=1, column=j).column_letter].width = wdt
        return start_row + len(data) + 2

    def stat_rows(sm):
        out = []
        for L, ph, s in METRIC_ORDER:
            for v in SHOWN:
                st = sm[(L, ph, s, v)]
                ann = ANN_CELLS.get((L, ph, s)) if v == "signups_added" else None
                out.append(
                    [
                        f"{L}",
                        PHASE_LABEL[ph],
                        SETTING_LABEL[s],
                        VALUE_LABEL[v],
                        st["mean"],
                        st["sd"],
                        st["p5"],
                        st["p50"],
                        st["p95"],
                        ann if ann is not None else "",
                        round(st["mean"] - ann, 2) if ann is not None else "",
                    ]
                )
        for c in CHOICES:
            for v in ("asked", "completed", "stopped_responding"):
                st = sm[("asking", c, v)]
                out.append(
                    [
                        "-",
                        f"Asking for more: {c}",
                        "-",
                        v.replace("_", " "),
                        st["mean"],
                        st["sd"],
                        st["p5"],
                        st["p50"],
                        st["p95"],
                        "",
                        "",
                    ]
                )
        return out

    hdr = [
        "Emails",
        "Event / step",
        "Setting",
        "Measure",
        "Mean",
        "SD",
        "5th pct",
        "Median",
        "95th pct",
        "Ann's number",
        "Ours minus Ann",
    ]
    wid = [8, 26, 20, 32, 9, 9, 9, 9, 9, 12, 12]

    ws = wb.active
    ws.title = "Summary"
    ws["A1"], ws["A1"].font = (
        "Class-exercise simulation: summary",
        Font(name="Arial", bold=True, size=14),
    )
    ws["A2"] = (
        f"{meta['rounds']} rounds x {meta['runs']:,} runs, seeds {meta['seed_ranges']}. 'Sign-ups added' = "
        "extra sign-ups the list brings over the mass email alone, same run. SD = spread across runs; "
        "5th-95th pct = the middle 90% of runs."
    )
    ws["A2"].font = base
    r = 4
    ws.cell(
        row=r, column=1, value="Balanced 3/3/2/2, 60 emails: sign-ups added, mean per round"
    ).font = bold
    rws = [
        [f"Round {i}"] + [sm[(60, ph, "balanced", "signups_added")]["mean"] for ph in PHASES]
        for i, sm in enumerate(summaries, 1)
    ]
    rws.append(
        ["All rounds"] + [combined[(60, ph, "balanced", "signups_added")]["mean"] for ph in PHASES]
    )
    rws.append(["Ann's number"] + [ANN_CELLS.get((60, ph, "balanced"), "") for ph in PHASES])
    r = table(ws, r + 1, ["Round"] + [PHASE_LABEL[ph] for ph in PHASES], rws)
    ws.cell(
        row=r, column=1, value="Every measure, every round (mean / SD / 5th-95th percentile)"
    ).font = bold
    long = []
    for L, ph, s in METRIC_ORDER:
        for v in SHOWN:
            row = [f"{L}", PHASE_LABEL[ph], SETTING_LABEL[s], VALUE_LABEL[v]]
            for sm in summaries + [combined]:
                st = sm[(L, ph, s, v)]
                row += [st["mean"], st["sd"], f"{st['p5']:g} to {st['p95']:g}"]
            ann = ANN_CELLS.get((L, ph, s)) if v == "signups_added" else None
            row.append(ann if ann is not None else "")
            long.append(row)
    names = [f"Round {i}" for i in range(1, len(summaries) + 1)] + ["All rounds"]
    sub = [f"{n} {x}" for n in names for x in ("mean", "SD", "5th-95th")]
    table(
        ws,
        r + 1,
        ["Emails", "Event / step", "Setting", "Measure"] + sub + ["Ann's number"],
        long,
        [8, 26, 20, 32] + [9, 9, 12] * len(names) + [12],
    )
    ws.freeze_panes = "E6"

    for i, sm in enumerate(summaries, 1):
        w = wb.create_sheet(f"Round {i}")
        w["A1"], w["A1"].font = (
            f"Round {i}: seeds {meta['round_seeds'][i - 1]}",
            Font(name="Arial", bold=True, size=13),
        )
        table(w, 3, hdr, stat_rows(sm), wid)
        w.freeze_panes = "E4"

    w = wb.create_sheet("Requirements check")
    counts = {s: sum(1 for x in rows if x["status"] == s) for s in ("PASS", "FAIL", "N/A")}
    w["A1"] = (
        f"{counts['PASS']} PASS, {counts['FAIL']} FAIL, {counts['N/A']} not checkable by simulation. "
        f"Ann's per-cell numbers pass when within {CELL_TOLERANCE} sign-ups."
    )
    w["A1"].font = bold
    r = table(
        w,
        3,
        [
            "ID",
            "Rule",
            "Ann's words",
            "Where it came from",
            "Target",
            "Our value",
            "Result",
            "Note",
        ],
        [
            [x[k] for k in ("id", "rule", "quote", "source", "target", "value", "status", "note")]
            for x in rows
        ],
        [11, 38, 48, 30, 26, 30, 9, 50],
    )
    w.cell(row=r, column=1, value="Conflicts between sources (newest from Ann wins)").font = bold
    table(
        w,
        r + 1,
        ["ID", "Topic", "Older", "Newest from Ann", "What we used"],
        [list(c) for c in CONFLICTS],
    )
    w.freeze_panes = "A4"

    w = wb.create_sheet("How to re-run")
    steps = [
        ["1", "Install Python 3.10 or newer, then run: pip install openpyxl"],
        ["2", "Put ann_simulation.py in a folder and open a terminal there."],
        ["3", "Quick check (about a minute): python ann_simulation.py --self-check"],
        ["4", "Full run: python ann_simulation.py --runs 5000 --rounds 3 --seed 1 --out results"],
        ["5", "Open results/ann-sim.xlsx. The same command always gives the same numbers."],
        ["", ""],
        ["Option", "Meaning"],
        ["--runs N", "Runs (simulated teams) per round. Default 5000."],
        [
            "--rounds N",
            "Independent rounds. Default 3. Round r uses seeds seed+(r-1)*runs to seed+r*runs-1.",
        ],
        [
            "--seed N",
            "First seed. Default 1. A different seed gives a different, equally valid set of runs.",
        ],
        ["--workers N", "Processor cores to use. Default: all. Does not change the numbers."],
        [
            "--out DIR",
            "Folder for ann-sim.xlsx, the summary JSON, one CSV of every run per round, and the report.",
        ],
        ["", ""],
        [
            "To change a rule",
            "Edit the numbers at the top of ann_simulation.py (for example ATTEND = 0.65) and re-run. "
            "Every assumption and its source is listed at the top of that file.",
        ],
        [
            "Asking your Claude",
            "'Run ann_simulation.py --self-check, then the full run, and tell me which checks fail.'",
        ],
    ]
    table(w, 1, ["Step", "What to do"], steps, [18, 110])
    wb.save(path)


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
def run_all(runs: int, rounds: int, seed: int, workers: int):
    seed_lists = [list(range(seed + i * runs, seed + (i + 1) * runs)) for i in range(rounds)]
    results = [run_round(s, workers) for s in seed_lists]
    return seed_lists, results


def self_check() -> None:
    """Small checks that fail loudly if the rule breaks."""
    assert len(PROFILES) == 300 and len(NO_CARD) == 230, "class file not loaded"
    assert one_run(1) == one_run(1), "same seed must give the same run"
    assert one_run(1) != one_run(2), "different seeds must differ"
    r = one_run(3)
    for L in LIMITS:
        for s in SETTINGS:
            for ph in PHASES:
                v = r[(L, ph, s)]
                assert v[0] <= v[1] and v[3] <= v[4] and v[4] <= v[1] and v[2] >= 0, (L, ph, s, v)
    assert r[("asking", "promise")] == (230, 69, 0) and r[("asking", "required")] == (230, 184, 35)
    lst = rank(HARBOR, "balanced", 60)
    assert len(lst) == 60 and lst == list_before(HARBOR, "balanced", 60)
    assert max(chances(HARBOR).values()) <= 0.70 + 1e-9
    seeds, res = run_all(200, 2, 1, 1)
    sm = summarise(res[0] + res[1])
    exp, got = (
        expected_added(NORTHLINE, "balanced", 60),
        sm[(60, "nl", "balanced", "signups_added")],
    )
    assert abs(got["mean"] - exp) < 5 * got["sd"] / 400**0.5, (got["mean"], exp)
    print(
        f"self-check OK: balanced Northline 60 added {got['mean']:.2f} over 400 runs (expected {exp:.2f})"
    )


def main(argv=None) -> None:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--runs", type=int, default=5000)
    ap.add_argument("--rounds", type=int, default=3)
    ap.add_argument("--seed", type=int, default=1)
    ap.add_argument("--workers", type=int, default=os.cpu_count() or 1)
    ap.add_argument("--out", default="results")
    ap.add_argument("--self-check", action="store_true")
    a = ap.parse_args(argv)
    if a.runs < 2 or a.rounds < 1:
        ap.error("--runs must be at least 2 and --rounds at least 1")
    if a.self_check:
        self_check()
        return
    self_check()
    t0 = time.time()
    seed_lists, rounds = run_all(a.runs, a.rounds, a.seed, a.workers)
    summaries = [summarise(rd) for rd in rounds]
    combined = summarise([r for rd in rounds for r in rd])
    det_ok = one_run(seed_lists[0][0]) == rounds[0][0]
    rows = requirement_rows(combined, rounds, det_ok)
    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)
    meta = {
        "date": time.strftime("%Y-%m-%d"),
        "runs": a.runs,
        "rounds": a.rounds,
        "seed": a.seed,
        "round_seeds": [f"{s[0]}-{s[-1]}" for s in seed_lists],
        "seed_ranges": ", ".join(f"{s[0]}-{s[-1]}" for s in seed_lists),
        "checksum": DATA["checksum"],
        "runtime_s": round(time.time() - t0),
    }
    for i, (s, rd) in enumerate(zip(seed_lists, rounds), 1):
        write_csv(out / f"runs-round{i}.csv", s, rd)
    key = lambda k: "|".join(map(str, k))  # noqa: E731
    (out / "summary.json").write_text(
        json.dumps(
            {
                "meta": meta,
                "rounds": [{key(k): v for k, v in sm.items()} for sm in summaries],
                "combined": {key(k): v for k, v in combined.items()},
                "requirements": rows,
            },
            indent=1,
        )
    )
    write_markdown(out / "requirements-check.md", rows, summaries, combined, meta)
    try:
        write_xlsx(out / "ann-sim.xlsx", rows, summaries, combined, meta)
    except ImportError:
        print("openpyxl not installed: skipped the spreadsheet (pip install openpyxl)")
    fails = [r["id"] for r in rows if r["status"] == "FAIL"]
    print(
        f"done in {time.time() - t0:.0f}s -> {out}/  PASS {sum(r['status'] == 'PASS' for r in rows)}, "
        f"FAIL {len(fails)} {fails}"
    )


if __name__ == "__main__":
    main()
