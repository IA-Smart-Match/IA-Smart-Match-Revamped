"""Open an uploaded Excel workbook safely, and hand back named sheets as text.

Bytes in; sheets of trimmed text cells, or one plain sentence, out. Nothing
here knows a column name — :mod:`smartmatch_domain.exercise.ingest` reads the
cells — and nothing here opens a path: the workbook is read from memory.

The owner ruled on 2026-09-24 (OQ-CE-05) that the instructor uploads Ann's
``.xlsx`` as-is, read with ``openpyxl`` in ``read_only`` and ``data_only``
mode, with every guard the CSV reader had plus a zip-bomb guard.

Untrusted input, in the order it is checked
===========================================
1. **Bytes.** :data:`MAX_UPLOAD_BYTES` before anything is read.
2. **Container.** An ``.xlsx`` is a ZIP. The first four bytes must say so; an
   old binary ``.xls`` and a CSV are each refused with their own sentence.
3. **Zip bomb, before openpyxl opens the file.** The ZIP central directory is
   read with :mod:`zipfile` and refused when it holds more than
   :data:`MAX_ZIP_ENTRIES` parts, when any part or all parts together would
   decompress past :data:`MAX_ENTRY_UNCOMPRESSED_BYTES` or
   :data:`MAX_TOTAL_UNCOMPRESSED_BYTES`, when a part is encrypted, or when a
   part uses a compression method other than stored or deflate. The declared
   sizes are a real bound and not only a claim: CPython's ``ZipExtFile``
   stops returning bytes at the declared ``file_size`` and raises on a CRC
   mismatch, so a part that lies about its size yields an error rather than
   more data.
4. **XML.** openpyxl parses with ``defusedxml`` when it is importable, and
   this module refuses every upload when openpyxl reports that it is not
   (:func:`_xml_is_defused`) — a missing dependency fails closed rather than
   parsing with the standard library's entity handling.
5. **Shape.** Rows come through :func:`itertools.islice` one past
   :data:`MAX_DATA_ROW_COUNT`, columns are read only up to one past
   :data:`MAX_COLUMN_COUNT`, and a cell longer than
   :data:`MAX_CELL_CHARACTERS` is refused.

Formulas are never evaluated: ``data_only=True`` returns the value Excel
cached, and a leading ``=`` in a text cell is text. External links are not
loaded (``keep_links=False``).

Any exception openpyxl raises while reading is turned into one sentence, and
only the exception's **class name** reaches the log — never its text, which
can quote the file.
"""

from __future__ import annotations

import datetime as dt
import io
import itertools
import logging
import warnings
import zipfile
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from typing import Final

import openpyxl
from openpyxl.reader.excel import ExcelReader
from openpyxl.styles.stylesheet import apply_stylesheet
from openpyxl.worksheet._read_only import ReadOnlyWorksheet

from smartmatch_domain.exercise.layout import IngestRefusal
from smartmatch_domain.ingest import normalize_header

__all__ = [
    "MAX_CELL_CHARACTERS",
    "MAX_COLUMN_COUNT",
    "MAX_DATA_ROW_COUNT",
    "MAX_ENTRY_UNCOMPRESSED_BYTES",
    "MAX_SHEET_COUNT",
    "MAX_TOTAL_UNCOMPRESSED_BYTES",
    "MAX_UPLOAD_BYTES",
    "MAX_ZIP_ENTRIES",
    "XLSX_MEDIA_TYPE",
    "SheetRows",
    "quote",
    "read_sheets",
]

_LOGGER = logging.getLogger(__name__)

# openpyxl reports odd workbook parts through ``warnings.warn``, and its
# messages quote the file (a defined name, a property name, a part path). That
# would reach stderr around the "class name only" logging rule, and every
# distinct message is kept forever in the module's ``__warningregistry__``.
# Silenced once, at import: ``warnings.catch_warnings`` is process-global and
# not safe in the threadpool the upload route runs in. Security review of
# PR #228, finding 3.
warnings.filterwarnings("ignore", module=r"openpyxl(\..*)?$")

#: The largest upload this will look at, in bytes. Ann's 300-profile workbook
#: is about 40 KB; two mebibytes leaves room without leaving room for a file
#: that is not a class roster at all. Checked before anything is read.
MAX_UPLOAD_BYTES: Final[int] = 2 * 1024 * 1024

#: The most parts the ZIP may hold. Ann's workbook has 18.
MAX_ZIP_ENTRIES: Final[int] = 100

#: The most one part may decompress to. Ann's largest part is about 180 KB and
#: a 1000-profile sheet about 0.6 MB. Kept tight because openpyxl's style and
#: property parsers build an object per element: 8 MB of styles cost a
#: gigabyte of memory in the security review of PR #228.
MAX_ENTRY_UNCOMPRESSED_BYTES: Final[int] = 1024 * 1024

#: The most all parts together may decompress to. Ann's workbook is ~300 KB.
MAX_TOTAL_UNCOMPRESSED_BYTES: Final[int] = 4 * 1024 * 1024

#: The most sheets the workbook may list. Ann's have 3 and 4. openpyxl sizes
#: every listed sheet by parsing it, and many ``<sheet>`` entries can point at
#: one part — so this cap, and the distinct-part rule beside it, bound how many
#: times the same bytes are parsed.
MAX_SHEET_COUNT: Final[int] = 16

#: The longest a single cell may be, in characters.
MAX_CELL_CHARACTERS: Final[int] = 500

#: The most columns a sheet's heading row may carry. Ann's widest has 14.
MAX_COLUMN_COUNT: Final[int] = 64

#: The most data rows one sheet may carry. Ann's file has 300 and 12.
MAX_DATA_ROW_COUNT: Final[int] = 4_000

#: The media type of an ``.xlsx`` upload, as the browser and the route name it.
XLSX_MEDIA_TYPE: Final[str] = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

#: What every ZIP starts with, and what an old binary ``.xls`` starts with.
_ZIP_MAGIC: Final[bytes] = b"PK\x03\x04"
_OLE_MAGIC: Final[bytes] = b"\xd0\xcf\x11\xe0"

#: ``zipfile`` compression methods a workbook may use: stored and deflate.
_ALLOWED_COMPRESSION: Final[frozenset[int]] = frozenset({zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED})

#: How much of a heading a sentence shows.
_QUOTED_TEXT_CHARACTERS: Final[int] = 40

_UNREADABLE: Final[str] = (
    "The file could not be read as an Excel workbook; please open it in Excel, "
    "save it as .xlsx and upload it again."
)


@dataclass(frozen=True, slots=True)
class SheetRows:
    """One sheet, once read: its headings and its non-blank data rows.

    Attributes:
        title: The sheet's name as the workbook spells it.
        headers: Normalized heading to the heading as the workbook spells it.
        rows: Each non-blank data row as ``(row number, {heading: text})``,
            where the row number is the one Excel shows.
    """

    title: str
    headers: Mapping[str, str]
    rows: tuple[tuple[int, Mapping[str, str]], ...]


def quote(text: str) -> str:
    """Render a piece of the uploaded file for a sentence an instructor reads.

    File content can carry newlines, control characters, a backtick that
    breaks out of the quoting these sentences use, and ten thousand characters
    where forty would do. Unprintables become spaces, whitespace collapses,
    backticks are removed, and the result is truncated. Empty reads as
    ``(blank)``. Presentation only: nothing here is ever stored.
    """
    flattened = "".join(character if character.isprintable() else " " for character in text)
    cleaned = " ".join(flattened.replace("`", "").split())
    if not cleaned:
        return "(blank)"
    if len(cleaned) > _QUOTED_TEXT_CHARACTERS:
        return cleaned[:_QUOTED_TEXT_CHARACTERS] + "…"
    return cleaned


def read_sheets(raw: bytes, sheet_names: Sequence[str]) -> tuple[SheetRows, ...] | IngestRefusal:
    """Read the named sheets out of an uploaded workbook, in the order named.

    Args:
        raw: The uploaded bytes.
        sheet_names: The sheets to read. A sheet is found by its normalized
            name, so ``"profiles"`` finds ``Profiles``. Every other sheet is
            never iterated.

    Returns:
        One :class:`SheetRows` per name, or the first refusal.
    """
    refusal = _guard_bytes(raw) or _guard_zip(raw)
    if refusal is not None:
        return refusal
    if not _xml_is_defused():
        _LOGGER.error("exercise workbook refused: openpyxl is not using defusedxml")
        return IngestRefusal(
            "workbook_reader_unavailable",
            "Workbooks cannot be read on this server right now; please tell the "
            "person who runs it.",
        )
    try:
        workbook = _open_workbook(raw)
    except Exception as error:  # untrusted-parser boundary; see the module docstring
        _LOGGER.info("exercise workbook unreadable: %s", type(error).__name__)
        return IngestRefusal("unreadable_workbook", _UNREADABLE)
    if isinstance(workbook, IngestRefusal):
        return workbook
    try:
        return _read_named(workbook, sheet_names)
    except Exception as error:  # untrusted-parser boundary; see the module docstring
        _LOGGER.info("exercise workbook unreadable: %s", type(error).__name__)
        return IngestRefusal("unreadable_workbook", _UNREADABLE)
    finally:
        workbook.close()


def _xml_is_defused() -> bool:
    """Whether openpyxl will parse all of this workbook's XML with ``defusedxml``.

    With lxml installed, openpyxl routes ``fromstring`` through lxml instead of
    defusedxml, so lxml's presence refuses too: the claim is "defusedxml
    everywhere", and it is checked rather than assumed.
    """
    return getattr(openpyxl, "DEFUSEDXML", False) is True and not getattr(openpyxl, "LXML", False)


def _open_workbook(raw: bytes) -> openpyxl.Workbook | IngestRefusal:
    """``openpyxl.load_workbook(read_only=True, data_only=True)``, one step at a time.

    The same steps as ``ExcelReader.read``, with one check between reading the
    workbook part and sizing its sheets: at most :data:`MAX_SHEET_COUNT`
    sheets, each on a part of its own. Without it, one sheet part listed a
    hundred thousand times is parsed a hundred thousand times — a 1.5 MB upload
    that passed every size guard and kept a CPU busy for hours in the security
    review of PR #228 (finding 1). Checked on openpyxl's own parse of the
    package rather than on a pre-read of ``xl/workbook.xml``, because
    ``[Content_Types].xml`` decides which part is the workbook.
    """
    reader = ExcelReader(io.BytesIO(raw), read_only=True, data_only=True, keep_links=False)
    try:
        reader.read_manifest()
        reader.read_strings()
        reader.read_workbook()
        targets = [rel.target for _, rel in reader.parser.find_sheets()]
        if len(targets) > MAX_SHEET_COUNT or len(set(targets)) != len(targets):
            reader.archive.close()
            return IngestRefusal(
                "too_many_sheets",
                f"The workbook has more than {MAX_SHEET_COUNT} sheets or sheets that "
                "share their contents; please upload the class data file itself.",
            )
        reader.read_properties()
        reader.read_custom()
        reader.read_theme()
        apply_stylesheet(reader.archive, reader.wb)
        reader.read_worksheets()
        reader.parser.assign_names()
    except BaseException:
        reader.archive.close()
        raise
    return reader.wb


# ---------------------------------------------------------------------------
# Guards that run before openpyxl sees a byte
# ---------------------------------------------------------------------------


def _guard_bytes(raw: bytes) -> IngestRefusal | None:
    """Size first, then "is this a workbook at all"."""
    if len(raw) > MAX_UPLOAD_BYTES:
        return IngestRefusal(
            "file_too_large",
            f"The file is larger than {MAX_UPLOAD_BYTES // (1024 * 1024)} MB, "
            "which is much larger than a class data file; please check you "
            "picked the right file.",
        )
    if not raw.strip():
        return IngestRefusal("file_empty", "The file is empty.")
    if raw.startswith(_OLE_MAGIC):
        return IngestRefusal(
            "old_excel_format",
            "This is an older Excel file (.xls); please open it in Excel, save it "
            "as an Excel Workbook (.xlsx) and upload that file instead.",
        )
    if not raw.startswith(_ZIP_MAGIC):
        return IngestRefusal(
            "not_a_workbook",
            "This page reads the class data file as an Excel workbook (.xlsx); "
            "please upload the .xlsx file itself rather than a CSV or other export.",
        )
    return None


def _guard_zip(raw: bytes) -> IngestRefusal | None:
    """The zip-bomb guard: the central directory, read before anything inflates."""
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            entries = archive.infolist()
    except Exception as error:  # untrusted-parser boundary; zipfile raises NotImplementedError too
        _LOGGER.info("exercise workbook unreadable: %s", type(error).__name__)
        return IngestRefusal("unreadable_workbook", _UNREADABLE)
    too_big = IngestRefusal(
        "workbook_too_large",
        "The workbook unpacks to much more than a class data file would; please "
        "check you picked the right file.",
    )
    if len(entries) > MAX_ZIP_ENTRIES:
        return too_big
    if any(entry.file_size > MAX_ENTRY_UNCOMPRESSED_BYTES for entry in entries):
        return too_big
    if sum(entry.file_size for entry in entries) > MAX_TOTAL_UNCOMPRESSED_BYTES:
        return too_big
    if any(entry.flag_bits & 0x1 for entry in entries):
        return IngestRefusal(
            "workbook_encrypted",
            "The workbook is password-protected; please remove the password in "
            "Excel and upload it again.",
        )
    if any(entry.compress_type not in _ALLOWED_COMPRESSION for entry in entries):
        return IngestRefusal("unreadable_workbook", _UNREADABLE)
    return None


# ---------------------------------------------------------------------------
# Reading the sheets
# ---------------------------------------------------------------------------


def _read_named(
    workbook: openpyxl.Workbook, sheet_names: Sequence[str]
) -> tuple[SheetRows, ...] | IngestRefusal:
    """Find each named sheet, then read it. The first failure is the answer."""
    by_name: dict[str, list[str]] = {}
    for name in workbook.sheetnames:
        by_name.setdefault(normalize_header(name), []).append(name)
    read: list[SheetRows] = []
    for wanted in sheet_names:
        titles = by_name.get(normalize_header(wanted), [])
        if not titles:
            return IngestRefusal("missing_sheet", f"The workbook has no sheet named `{wanted}`.")
        if len(titles) > 1:
            return IngestRefusal(
                "duplicate_sheet",
                f"The workbook has more than one sheet named like `{wanted}`; please "
                "leave one of them and upload it again.",
            )
        title = titles[0]
        worksheet = workbook[title]
        if not isinstance(worksheet, ReadOnlyWorksheet):
            return IngestRefusal(
                "missing_sheet", f"The workbook's `{wanted}` sheet is not a table of cells."
            )
        # Sentences name the sheet the layout asked for, not the workbook's own
        # title: a title only has to *normalize* to ``profiles``, so it can carry
        # backticks, newlines or two million characters (security review, 4).
        sheet = _read_sheet(wanted, worksheet)
        if isinstance(sheet, IngestRefusal):
            return sheet
        read.append(sheet)
    return tuple(read)


def _read_sheet(title: str, worksheet: ReadOnlyWorksheet) -> SheetRows | IngestRefusal:
    """The heading row and the data rows of one sheet, bounded.

    The heading row is the first non-blank row among the first
    :data:`MAX_DATA_ROW_COUNT` — bounded too, because a sheet whose first cell
    sits a million rows down would otherwise be walked a row at a time. A data
    row past the cap is what the cap refuses on, which is why the sentence
    says "more than".
    """
    # types-openpyxl declares ``ReadOnlyWorksheet.iter_rows`` as a copy of
    # ``Worksheet.iter_rows``, self type included; the runtime method is the
    # read-only one, which streams.
    # A read-only sheet stops at the row its ``<dimension>`` declares, and a
    # stale declaration would drop rows without a word (ADR-0011). Reset it so
    # the sheet is read to its real end; the ``islice`` caps bound the work.
    worksheet.reset_dimensions()
    raw_rows = worksheet.iter_rows(max_col=MAX_COLUMN_COUNT + 1, values_only=True)  # type: ignore[misc]
    numbered = enumerate(raw_rows, start=1)
    header: tuple[int, list[str]] | None = None
    for number, values in itertools.islice(numbered, MAX_DATA_ROW_COUNT):
        cells = [_cell_text(value) for value in values]
        if any(cells):
            header = (number, cells)
            break
    if header is None:
        return IngestRefusal(
            "no_header_row", f"The `{title}` sheet has no heading row naming its columns."
        )
    headers = _header_map(title, header[1])
    if isinstance(headers, IngestRefusal):
        return headers
    body = list(itertools.islice(numbered, MAX_DATA_ROW_COUNT + 1))
    if len(body) > MAX_DATA_ROW_COUNT:
        return IngestRefusal(
            "too_many_rows",
            f"The `{title}` sheet has more than {MAX_DATA_ROW_COUNT} rows, which is "
            "more than this page reads.",
        )
    return _collect_rows(title, header[1], headers, body)


def _header_map(title: str, cells: Sequence[str]) -> Mapping[str, str] | IngestRefusal:
    """Normalized heading to the sheet's own spelling, refusing what is unusable."""
    if len(cells) > MAX_COLUMN_COUNT and cells[MAX_COLUMN_COUNT]:
        return IngestRefusal(
            "too_many_columns",
            f"The `{title}` sheet has more than {MAX_COLUMN_COUNT} columns, which is "
            "more than this page reads.",
        )
    mapped: dict[str, str] = {}
    for heading in cells[:MAX_COLUMN_COUNT]:
        if len(heading) > MAX_CELL_CHARACTERS:
            return IngestRefusal(
                "column_name_too_long",
                f"One of the column headings on the `{title}` sheet is longer than "
                f"{MAX_CELL_CHARACTERS} characters; please check its first row.",
            )
        key = normalize_header(heading)
        if key and key in mapped:
            return IngestRefusal(
                "duplicate_column",
                f"The `{title}` sheet has two columns named `{quote(heading)}`; please "
                "leave one of them and upload it again.",
            )
        if key:
            mapped[key] = heading
    return mapped


def _collect_rows(
    title: str,
    heading_cells: Sequence[str],
    headers: Mapping[str, str],
    body: Sequence[tuple[int, Sequence[object]]],
) -> SheetRows | IngestRefusal:
    """Bound every cell, key it by its heading, and drop entirely blank rows."""
    kept: list[tuple[int, Mapping[str, str]]] = []
    for number, values in body:
        cells = [_cell_text(value) for value in values[:MAX_COLUMN_COUNT]]
        row: dict[str, str] = {}
        for index, text in enumerate(cells):
            heading = heading_cells[index] if index < len(heading_cells) else ""
            if not heading:
                if text:
                    return IngestRefusal(
                        "cell_without_heading",
                        f"Row {number} of the `{title}` sheet has a value in a column "
                        "with no heading; please give the column a heading or clear it.",
                    )
                continue
            if len(text) > MAX_CELL_CHARACTERS:
                return IngestRefusal(
                    "cell_too_long",
                    f"Row {number} of the `{title}` sheet has more than "
                    f"{MAX_CELL_CHARACTERS} characters in the column "
                    f"`{quote(heading)}`; please shorten it and upload again.",
                )
            row[heading] = text
        if any(row.values()):
            kept.append((number, row))
    return SheetRows(title=title, headers=headers, rows=tuple(kept))


def _cell_text(value: object) -> str:
    """One cell as trimmed text. A missing cell and a blank one are both blank.

    A whole number typed as a float (``186.0``) reads as ``186``, because
    Excel stores every number as a float and a profile's place in the fixed
    order is a whole number however it was typed. A date reads as its ISO
    form. Nothing is evaluated.
    """
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, dt.datetime | dt.date | dt.time):
        return value.isoformat()
    return str(value).strip()
