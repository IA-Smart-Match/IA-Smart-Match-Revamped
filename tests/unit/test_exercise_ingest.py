"""The class-exercise data-file parser: what it accepts, and what it says.

The file under test is ``smartmatch_domain/exercise/ingest.py``, reading Ann's
``.xlsx`` (OQ-CE-01 and OQ-CE-05, closed 2026-09-24). These tests assert:

* that Ann's own 300-profile workbook parses, and what it parses into;
* that each check of design spec §3 happens and produces one sentence an
  instructor can act on — a missing sheet or column names the sheet;
* that the closed vocabularies refuse what is not in them;
* that ADR-0025 D6's two withheld columns reach :class:`ParsedProfile` and
  nothing else — asserted by walking a corpus of bad files as well as good
  ones, because a leak through an error message is the leak nobody looks for;
* that the checksum is a checksum of the uploaded bytes and does not move.

The workbook-level guards (byte cap, zip bomb, XML) are
``test_exercise_workbook.py``'s.
"""

from __future__ import annotations

import dataclasses
import hashlib
import logging

import pytest
from smartmatch_domain.exercise.ingest import (
    EXERCISE_EVENT_ROW_COUNT,
    MAX_COLUMN_INTEGER,
    MIN_PROFILE_ROW_COUNT,
    ExerciseFileLayout,
    IngestRefusal,
    ParsedDataset,
    parse_exercise_file,
)
from smartmatch_domain.exercise.vocabulary import EXERCISE_MAJORS
from smartmatch_domain.student_factors import EventEvidence, ProfileEvidence, same_major

from tests.unit.exercise_workbooks import (
    ANN_FULL_FILE,
    ANN_SAMPLE_FILE,
    EVENT_HEADINGS,
    LAYOUT,
    PROFILE_HEADINGS,
    WITHHELD_GOAL,
    WITHHELD_TOPIC,
    event_rows,
    good_profile_rows,
    good_workbook,
    workbook_bytes,
)


def _accepted(content: bytes, layout: ExerciseFileLayout = LAYOUT) -> ParsedDataset:
    result = parse_exercise_file(content, layout=layout)
    assert isinstance(result, ParsedDataset), result
    return result


def _refusal(content: bytes) -> IngestRefusal:
    result = parse_exercise_file(content)
    assert isinstance(result, IngestRefusal), result
    return result


def _with_profile(index: int, **overrides: object) -> bytes:
    rows = good_profile_rows()
    rows[index].update(overrides)
    return workbook_bytes(rows, event_rows())


def _with_event(index: int, **overrides: object) -> bytes:
    events = event_rows()
    events[index].update(overrides)
    return workbook_bytes(good_profile_rows(), events)


# ---------------------------------------------------------------------------
# Ann's own file
# ---------------------------------------------------------------------------


def test_anns_full_workbook_is_accepted_as_she_sent_it() -> None:
    dataset = _accepted(ANN_FULL_FILE.read_bytes())

    assert dataset.row_count == 300
    assert len(dataset.events) == 12
    assert [event.event_key for event in dataset.events if event.is_exercise_event] == [
        "E11",
        "E12",
    ]
    assert dataset.report.distinct_class_years == ("Freshman", "Sophomore", "Junior", "Senior")
    assert dataset.report.profiles_without_card == 230
    assert dataset.report.distinct_stated_interest_terms == 13


def test_northline_and_harbor_are_exploratory_events_in_anns_file() -> None:
    """OQ-CE-14 (Ann, 2026-09-25): both are company events, so both are exploratory.

    Of the ten past events, the career fair, the three employer info sessions and
    the three industry panels are exploratory; the workshop, the competition and
    the networking mixer are not.
    """
    dataset = _accepted(ANN_FULL_FILE.read_bytes())
    exploratory = {event.event_key for event in dataset.events if event.is_exploratory}

    assert exploratory == {"E02", "E03", "E04", "E05", "E07", "E08", "E09", "E11", "E12"}


def test_the_markers_agree_with_anns_info_level_column() -> None:
    """Ann: info_level "can also be computed by the app". It is, and it agrees."""
    markers = _accepted(ANN_FULL_FILE.read_bytes()).report.markers

    assert (markers.major_only, markers.major_plus_events, markers.completed_card) == (166, 64, 70)


def test_p004_is_read_the_way_anns_read_me_describes_it() -> None:
    dataset = _accepted(ANN_FULL_FILE.read_bytes())
    p004 = next(profile for profile in dataset.profiles if profile.profile_no == 4)

    assert p004.display_name == "Brandon Soto"
    assert (p004.major, p004.class_year) == ("Accounting", "Senior")
    assert p004.past_event_keys == ("E07",)
    assert p004.stated_interests is not None
    assert p004.stated_interests[0] == "Technology / information systems"
    assert p004.career_goal == "Data, analytics or IT role"
    assert p004.tiebreak_order == 186
    assert p004.hidden_true_career_goal == "Data, analytics or IT role"


def test_the_twenty_row_sample_is_refused_by_the_specs_row_count_floor() -> None:
    """Design spec §3: 50–1000 profiles. The sample is for reading, not uploading."""
    refusal = _refusal(ANN_SAMPLE_FILE.read_bytes())

    assert refusal.code == "row_count_out_of_range"
    assert refusal.message == "The `Profiles` sheet has 20 profiles; it needs between 50 and 1000."


def test_the_sample_is_an_exact_subset_of_the_full_file() -> None:
    """Ann: "The sample is 20 rows taken from the full file (same IDs, same values)"."""
    from openpyxl import load_workbook

    def rows(path: object) -> dict[object, tuple[object, ...]]:
        book = load_workbook(path, read_only=True, data_only=True)  # type: ignore[arg-type]
        try:
            return {row[0]: row for row in book["Profiles"].iter_rows(min_row=2, values_only=True)}
        finally:
            book.close()

    sample, full = rows(ANN_SAMPLE_FILE), rows(ANN_FULL_FILE)
    assert len(sample) == 20
    assert all(full[key] == row for key, row in sample.items())


# ---------------------------------------------------------------------------
# A good file, and what it becomes
# ---------------------------------------------------------------------------


def test_a_well_formed_workbook_is_parsed_into_profiles_and_events() -> None:
    dataset = _accepted(good_workbook())

    assert dataset.row_count == 60
    assert dataset.report.exercise_event_count == EXERCISE_EVENT_ROW_COUNT
    assert [event.sequence for event in dataset.events] == list(range(1, 13))
    assert dataset.profiles[0].display_name == "Fictional Profile 001"


def test_a_profile_id_is_p_and_its_number() -> None:
    assert _accepted(good_workbook()).profiles[3].profile_no == 4


@pytest.mark.parametrize("profile_id", ["4", "Q004", "P-4", "P", "P0", "P4.5", "P٣", "P 4"])
def test_a_profile_id_that_is_not_p_and_digits_is_refused(profile_id: str) -> None:
    refusal = _refusal(_with_profile(0, profile_id=profile_id))

    assert refusal.code == "bad_profile_id"
    assert refusal.message.startswith("Row 2 of the `Profiles` sheet has `")
    assert "a profile id is P followed by a number, such as P004." in refusal.message


def test_a_profile_number_too_large_for_the_column_is_refused() -> None:
    refusal = _refusal(_with_profile(0, profile_id=f"P{MAX_COLUMN_INTEGER + 1}"))

    assert refusal.code == "bad_profile_id"


def test_no_card_is_none_and_a_card_with_no_interests_is_an_empty_card() -> None:
    dataset = _accepted(_with_profile(2, stated_interests=None, stated_career_goal=None))

    assert dataset.profiles[0].stated_interests is None  # card_completed = No
    assert dataset.profiles[2].stated_interests == ()  # card_completed = Yes, nothing listed
    assert dataset.profiles[2].career_goal is None


def test_a_card_that_says_no_but_carries_interests_is_refused() -> None:
    refusal = _refusal(_with_profile(0, stated_interests="Consulting"))

    assert refusal.code == "card_contradiction"
    assert refusal.message == (
        "Row 2 of the `Profiles` sheet says No in the column `card_completed` but has a "
        "value in the column `stated_interests`; a profile without a card has no stated "
        "interests and no stated career goal."
    )


def test_an_unreadable_card_flag_is_refused_rather_than_guessed() -> None:
    refusal = _refusal(_with_profile(0, card_completed="maybe"))

    assert refusal.code == "bad_yes_no"


def test_terms_are_stored_in_anns_spelling_whatever_the_cell_spelling() -> None:
    dataset = _accepted(
        _with_profile(
            2,
            stated_interests="  technology / INFORMATION systems ;consulting;;Consulting",
            major="accounting",
            year="SENIOR",
        )
    )

    profile = dataset.profiles[2]
    assert profile.stated_interests == ("Technology / information systems", "Consulting")
    assert (profile.major, profile.class_year) == ("Accounting", "Senior")


def test_all_majors_is_every_major_and_same_major_matches_it() -> None:
    """Brief: define "All majors" explicitly as match-all, and test it."""
    dataset = _accepted(good_workbook())
    everyone = dataset.events[0]
    assert everyone.target_majors == EXERCISE_MAJORS

    event = EventEvidence(event_key=everyone.event_key, target_majors=everyone.target_majors)
    for major in EXERCISE_MAJORS:
        assert same_major(ProfileEvidence(profile_id="1", major=major), event).value == 1.0


def test_a_single_target_major_is_that_major_only() -> None:
    dataset = _accepted(good_workbook())

    assert dataset.events[10].target_majors == ("Computer Information Systems",)


# ---------------------------------------------------------------------------
# Closed vocabularies (owner ruling 6, 2026-09-24)
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("column", "value", "ending"),
    [
        ("major", "Basket Weaving", "which is not one of the 6 majors."),
        ("year", "Fifth year", "which is not Freshman, Sophomore, Junior or Senior."),
        ("stated_career_goal", "Astronaut", "which is not one of the 16 career goals."),
    ],
)
def test_a_value_outside_the_vocabulary_is_refused_and_named(
    column: str, value: str, ending: str
) -> None:
    refusal = _refusal(_with_profile(2, **{column: value}))

    assert refusal.code == "unknown_value"
    assert refusal.message == (
        f"Row 4 of the `Profiles` sheet has `{value}` in the column `{column}`, {ending}"
    )


def test_an_unknown_topic_on_a_card_is_refused() -> None:
    refusal = _refusal(_with_profile(2, stated_interests="Consulting;Space tourism"))

    assert refusal.message == (
        "Row 4 of the `Profiles` sheet has `Space tourism` in the column "
        "`stated_interests`, which is not one of the 13 topics."
    )


def test_an_unknown_event_topic_or_target_major_is_refused() -> None:
    topic = _refusal(_with_event(0, event_topics="Underwater basket weaving"))
    major = _refusal(_with_event(3, target_major="Undeclared"))

    assert "Row 2 of the `Events` sheet has `Underwater basket weaving`" in topic.message
    assert major.message.endswith('which is not one of the 6 majors or "All majors".')


def test_an_unknown_event_type_is_refused_and_named() -> None:
    refusal = _refusal(_with_event(0, event_type="Hackathon"))

    assert refusal.code == "unknown_value"
    assert refusal.message.startswith("Row 2 of the `Events` sheet has `Hackathon`")
    assert refusal.message.endswith("which is not one of the 9 event types.")


def test_an_event_with_no_event_type_is_refused() -> None:
    refusal = _refusal(_with_event(10, event_type=None))

    assert refusal.message == "Row 12 of the `Events` sheet has nothing in the column `event_type`."


def test_an_event_type_is_read_through_the_fold() -> None:
    dataset = _accepted(_with_event(0, event_type="  career FAIR "))

    assert dataset.events[0].is_exploratory is True


@pytest.mark.parametrize("column", ["major", "year", "first_name", "last_name"])
def test_a_profile_missing_a_required_value_is_refused(column: str) -> None:
    refusal = _refusal(_with_profile(0, **{column: None}))

    assert refusal.message == f"Row 2 of the `Profiles` sheet has nothing in the column `{column}`."


# ---------------------------------------------------------------------------
# Sheets and columns
# ---------------------------------------------------------------------------


def test_a_missing_sheet_is_named() -> None:
    content = workbook_bytes(good_profile_rows(), event_rows(), sheet_names=("Profiles", "Evts"))

    assert _refusal(content).message == "The workbook has no sheet named `Events`."


def test_a_missing_column_names_the_sheet_and_the_column() -> None:
    headings = [h for h in PROFILE_HEADINGS if h != LAYOUT.tiebreak_order_column]
    content = workbook_bytes(good_profile_rows(), event_rows(), profile_headings=headings)

    refusal = _refusal(content)
    assert refusal.code == "missing_columns"
    assert refusal.message == "The `Profiles` sheet is missing the column `tiebreak_order`."


def test_every_missing_column_of_a_sheet_is_named_in_one_sentence() -> None:
    headings = [h for h in EVENT_HEADINGS if h not in {"seats", "event_topics"}]
    content = workbook_bytes(good_profile_rows(), event_rows(), event_headings=headings)

    assert _refusal(content).message == (
        "The `Events` sheet is missing these columns: `event_topics`, `seats`."
    )


def test_the_columns_the_parser_does_not_read_are_not_required() -> None:
    profile_headings = [
        h for h in PROFILE_HEADINGS if h not in {"events_attended_count", "info_level"}
    ]
    event_headings = [h for h in EVENT_HEADINGS if h != "event_date"]
    content = workbook_bytes(
        good_profile_rows(),
        event_rows(),
        profile_headings=profile_headings,
        event_headings=event_headings,
    )

    _accepted(content)


def test_heading_spelling_is_presentation_rather_than_identity() -> None:
    renamed = {"tiebreak_order": "Tiebreak Order", "card_completed": " CARD-completed "}
    headings = [renamed.get(h, h) for h in PROFILE_HEADINGS]
    rows = [{renamed.get(k, k): v for k, v in row.items()} for row in good_profile_rows()]

    _accepted(workbook_bytes(rows, event_rows(), profile_headings=headings))


def test_closing_a_rename_is_one_layout_object_and_no_parser_edit() -> None:
    """Every column name lives on the layout: a renamed file is a replaced layout."""
    layout = dataclasses.replace(LAYOUT, tiebreak_order_column="fixed_order", events_sheet="Talks")
    headings = [("fixed_order" if h == "tiebreak_order" else h) for h in PROFILE_HEADINGS]
    rows = [
        {("fixed_order" if k == "tiebreak_order" else k): v for k, v in row.items()}
        for row in good_profile_rows()
    ]
    content = workbook_bytes(
        rows, event_rows(), profile_headings=headings, sheet_names=("Profiles", "Talks")
    )

    assert _accepted(content, layout).profiles[0].tiebreak_order == 60


# ---------------------------------------------------------------------------
# Rows, counts and cross-row checks
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("profile_count", [MIN_PROFILE_ROW_COUNT - 1, 1001])
def test_a_file_outside_the_row_count_range_is_refused(profile_count: int) -> None:
    refusal = _refusal(good_workbook(profile_count))

    assert refusal.code == "row_count_out_of_range"


def test_one_exercise_event_is_not_enough() -> None:
    refusal = _refusal(_with_event(11, exercise_event="No"))

    assert refusal.code == "wrong_exercise_event_count"
    assert refusal.message == (
        "The file flags 1 events as exercise events; it needs exactly 2, one for each round."
    )


def test_three_exercise_events_are_too_many() -> None:
    refusal = _refusal(_with_event(0, exercise_event="Yes", seats=60))

    assert refusal.code == "wrong_exercise_event_count"
    assert "flags 3 events" in refusal.message


@pytest.mark.parametrize("seats", [None, 59, "sixty"])
def test_an_exercise_event_must_have_the_exercises_sixty_seats(seats: object) -> None:
    refusal = _refusal(_with_event(10, seats=seats))

    assert refusal.code == "bad_seats"
    assert refusal.message.startswith("Row 12 of the `Events` sheet gives `")
    assert refusal.message.endswith("this exercise runs with 60 seats per event.")


def test_a_past_events_seats_are_not_read() -> None:
    _accepted(_with_event(0, seats="n/a"))


def test_a_duplicate_profile_id_is_named_even_when_spelled_differently() -> None:
    refusal = _refusal(_with_profile(4, profile_id="P0004"))

    assert refusal.code == "duplicate_profile_id"
    assert refusal.message == "Two or more profiles share the same id: `P004`."


def test_a_duplicate_event_id_is_named() -> None:
    refusal = _refusal(_with_event(1, event_id="E01"))

    assert refusal.message == "Two or more events share the same id: `E01`."


def test_two_profiles_may_not_share_a_place_in_the_fixed_order() -> None:
    refusal = _refusal(_with_profile(1, tiebreak_order=60))

    assert refusal.code == "duplicate_tiebreak_order"
    assert "`tiebreak_order`: `60`" in refusal.message


@pytest.mark.parametrize("value", [None, 0, -3, 2.5, "first"])
def test_a_fixed_order_that_is_not_a_positive_whole_number_is_refused(value: object) -> None:
    refusal = _refusal(_with_profile(0, tiebreak_order=value))

    assert refusal.message == (
        "Row 2 of the `Profiles` sheet has no whole number in the column `tiebreak_order`."
    )


def test_a_fixed_order_typed_as_a_whole_float_is_read_as_the_number() -> None:
    assert _accepted(_with_profile(0, tiebreak_order=60.0)).profiles[0].tiebreak_order == 60


@pytest.mark.parametrize("key", ["E99", "E11"])
def test_an_attended_event_that_is_not_a_past_event_is_refused_not_dropped(key: str) -> None:
    refusal = _refusal(_with_profile(3, events_attended=f"E01;{key}"))

    assert refusal.code == "unknown_past_event_key"
    assert refusal.message == (
        "Some profiles list attended events that are not past events in the file: "
        f"`{key}` (1 rows)."
    )


def test_a_formula_is_never_evaluated() -> None:
    """``data_only``: a formula cell reads as Excel's cached value, never as a sum.

    openpyxl writes no cached value, so the cell reads as blank — and a blank
    name is refused, which is the proof that ``=1+1`` did not become ``2``.
    """
    refusal = _refusal(_with_profile(0, last_name="=1+1"))

    assert refusal.message == "Row 2 of the `Profiles` sheet has nothing in the column `last_name`."


# ---------------------------------------------------------------------------
# ADR-0025 D6 — the two withheld columns
# ---------------------------------------------------------------------------


def test_the_withheld_columns_reach_the_profile_and_nothing_else() -> None:
    dataset = _accepted(good_workbook())
    profile = dataset.profiles[0]

    assert profile.hidden_true_interests == (WITHHELD_TOPIC, "Consulting")
    assert profile.hidden_true_career_goal == WITHHELD_GOAL
    assert WITHHELD_TOPIC not in repr(dataset.report)
    assert not any("hidden" in field.name for field in dataclasses.fields(dataset.report))


def test_no_repr_in_the_parsed_result_prints_a_withheld_value() -> None:
    dataset = _accepted(good_workbook())
    profile = dataset.profiles[0]

    for rendered in (repr(profile), repr(dataset), repr(dataset.profiles), str(profile)):
        assert WITHHELD_TOPIC not in rendered
        assert "Healthcare" not in rendered
    assert "Fictional Profile 001" in repr(profile)


def _corpus_of_bad_files() -> list[bytes]:
    """One file per refusal path, each carrying the withheld values."""
    corpus = [
        _with_profile(0, first_name=None),
        _with_profile(0, profile_id="zero"),
        _with_profile(3, profile_id="P003"),
        _with_profile(2, events_attended="E77"),
        _with_profile(0, major="Nope"),
        _with_profile(0, card_completed="No", stated_career_goal="Undecided"),
        _with_profile(0, hidden_true_interests=f"{WITHHELD_TOPIC};Not a topic at all"),
        _with_profile(0, hidden_true_career_goal=f"{WITHHELD_GOAL} (secret)"),
        _with_profile(0, tiebreak_order="x"),
        _with_profile(0, last_name="y" * 900),
        _with_event(10, seats=1),
        good_workbook(3),
        workbook_bytes(
            good_profile_rows(),
            event_rows(),
            profile_headings=[h for h in PROFILE_HEADINGS if h != LAYOUT.major_column],
        ),
    ]
    return corpus


def test_no_refusal_sentence_ever_quotes_a_withheld_cell() -> None:
    for content in _corpus_of_bad_files():
        refusal = _refusal(content)
        assert "Healthcare" not in refusal.message, refusal
        assert "secret" not in refusal.message, refusal
        assert "Not a topic at all" not in refusal.message, refusal


def test_a_bad_withheld_cell_is_refused_by_naming_its_column_only() -> None:
    refusal = _refusal(_with_profile(0, hidden_true_career_goal="Astronaut"))

    assert refusal.message == (
        "Row 2 of the `Profiles` sheet has a value in the column `hidden_true_career_goal` "
        "that is not one of the 16 career goals."
    )


def test_the_log_records_counts_and_no_content(caplog: pytest.LogCaptureFixture) -> None:
    with caplog.at_level(logging.INFO, logger="smartmatch_domain.exercise"):
        _accepted(good_workbook())
        _refusal(good_workbook(3))
        _refusal(_with_profile(0, hidden_true_career_goal="Astronaut"))

    logged = "\n".join(record.getMessage() for record in caplog.records)
    assert "profiles=60" in logged
    assert "row_count_out_of_range" in logged
    assert "Healthcare" not in logged
    assert "Astronaut" not in logged
    assert "Fictional" not in logged


# ---------------------------------------------------------------------------
# The checksum
# ---------------------------------------------------------------------------


def test_the_checksum_is_stable_for_the_same_bytes() -> None:
    content = good_workbook()

    first, second = _accepted(content), _accepted(content)

    assert first.checksum == second.checksum == hashlib.sha256(content).hexdigest()


# ---------------------------------------------------------------------------
# Every column name lives on the layout
# ---------------------------------------------------------------------------


def test_every_layout_column_is_required_on_its_sheet() -> None:
    declared = {
        getattr(LAYOUT, field.name)
        for field in dataclasses.fields(LAYOUT)
        if field.name.endswith("_column")
    }

    assert set(LAYOUT.profile_columns) | set(LAYOUT.event_columns) == declared
    assert LAYOUT.withheld_columns == {"hidden_true_interests", "hidden_true_career_goal"}


def _column_names_written_inside_functions(module: object) -> list[tuple[str, str]]:
    """Every column name a *function* in ``module`` writes down as a literal.

    Walked with :mod:`ast` rather than searched as text, so a module's own
    prose is not mistaken for code that hard-codes a name. Docstrings are
    skipped; every other string literal inside a function is checked.
    """
    import ast
    from pathlib import Path

    names = {
        getattr(LAYOUT, field.name)
        for field in dataclasses.fields(LAYOUT)
        if field.name.endswith(("_column", "_sheet"))
    }
    source = Path(module.__file__).read_text(encoding="utf-8")  # type: ignore[attr-defined]
    tree = ast.parse(source)
    return [
        (node.name, literal.value)
        for node in ast.walk(tree)
        if isinstance(node, ast.FunctionDef)
        for statement in (node.body[1:] if ast.get_docstring(node) else node.body)
        for literal in ast.walk(statement)
        if isinstance(literal, ast.Constant)
        and isinstance(literal.value, str)
        and literal.value in names
    ]


@pytest.mark.parametrize("module_name", ["ingest", "layout", "workbook"])
def test_no_function_writes_a_column_or_sheet_name_down(module_name: str) -> None:
    import importlib

    module = importlib.import_module(f"smartmatch_domain.exercise.{module_name}")

    assert _column_names_written_inside_functions(module) == []


def test_the_column_name_guard_can_fail() -> None:
    """The assertion above is negative; prove the walk is not vacuous."""
    import types
    from pathlib import Path
    from tempfile import TemporaryDirectory

    with TemporaryDirectory() as directory:
        offender = Path(directory) / "offender.py"
        offender.write_text(
            '"""A docstring naming major, which must not count."""\n'
            "def read(row):\n"
            '    """Also naming major."""\n'
            f"    return row[{LAYOUT.major_column!r}]\n",
            encoding="utf-8",
        )
        module = types.ModuleType("offender")
        module.__file__ = str(offender)

        assert _column_names_written_inside_functions(module) == [("read", LAYOUT.major_column)]


# ---------------------------------------------------------------------------
# Review round 1
# ---------------------------------------------------------------------------


def test_all_majors_beside_something_that_is_not_a_major_is_refused() -> None:
    refusal = _refusal(_with_event(0, target_major="All majors; Basket Weaving"))

    assert refusal.message == (
        "Row 2 of the `Events` sheet has `Basket Weaving` in the column `target_major`, "
        'which is not one of the 6 majors or "All majors".'
    )


def test_an_event_attended_twice_in_one_cell_is_one_attendance() -> None:
    dataset = _accepted(_with_profile(3, events_attended="E01;E01; E01;E03"))

    assert dataset.profiles[3].past_event_keys == ("E01", "E03")
