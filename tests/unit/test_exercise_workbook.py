"""The workbook reader's guards: bytes, container, zip bomb, XML, shape.

The file under test is ``smartmatch_domain/exercise/workbook.py``. Every guard
the owner ruled on 2026-09-24 has a test here that a hostile or broken upload
is answered with one sentence, and that the zip-bomb guard answers **before**
openpyxl is asked to open anything.
"""

from __future__ import annotations

import io
import logging
import struct
import zipfile
from pathlib import Path

import openpyxl
import pytest
from smartmatch_domain.exercise import workbook
from smartmatch_domain.exercise.layout import IngestRefusal
from smartmatch_domain.exercise.workbook import (
    MAX_CELL_CHARACTERS,
    MAX_COLUMN_COUNT,
    MAX_DATA_ROW_COUNT,
    MAX_ENTRY_UNCOMPRESSED_BYTES,
    MAX_SHEET_COUNT,
    MAX_TOTAL_UNCOMPRESSED_BYTES,
    MAX_UPLOAD_BYTES,
    MAX_ZIP_ENTRIES,
    SheetRows,
    read_sheets,
)

from tests.unit.exercise_workbooks import good_workbook


def _sheet_bytes(*rows: tuple[object, ...], title: str = "Profiles") -> bytes:
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = title
    for row in rows:
        sheet.append(list(row))
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _read(raw: bytes, *names: str) -> tuple[SheetRows, ...] | IngestRefusal:
    return read_sheets(raw, names or ("Profiles",))


def _refused(raw: bytes, *names: str) -> IngestRefusal:
    result = _read(raw, *names)
    assert isinstance(result, IngestRefusal), result
    return result


def _zip(entries: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
    return buffer.getvalue()


def _replace_part(raw: bytes, name: str, data: bytes) -> bytes:
    """The same workbook with one part's bytes swapped."""
    source = zipfile.ZipFile(io.BytesIO(raw))
    return _zip(
        {
            info.filename: (data if info.filename == name else source.read(info))
            for info in source.infolist()
        }
    )


# ---------------------------------------------------------------------------
# Reading
# ---------------------------------------------------------------------------


def test_a_sheet_is_read_as_headings_and_trimmed_text_rows() -> None:
    raw = _sheet_bytes(
        ("id", "count", "flag"), (" P001 ", 186.0, True), (None, None, None), ("P002", 3, False)
    )

    result = _read(raw)
    assert isinstance(result, tuple)
    (sheet,) = result
    assert sheet.title == "Profiles"
    assert sheet.rows == (
        (2, {"id": "P001", "count": "186", "flag": "Yes"}),
        (4, {"id": "P002", "count": "3", "flag": "No"}),
    )


def test_a_sheet_is_found_by_its_normalized_name() -> None:
    result = _read(_sheet_bytes(("id",), ("P001",), title="profiles "), "Profiles")

    assert isinstance(result, tuple)


def test_the_heading_row_is_the_first_non_blank_row() -> None:
    raw = _sheet_bytes((None,), ("id",), ("P001",))

    result = _read(raw)
    assert isinstance(result, tuple)
    assert result[0].rows == ((3, {"id": "P001"}),)


def test_a_missing_sheet_is_named() -> None:
    assert _refused(_sheet_bytes(("id",)), "Profiles", "Events").message == (
        "The workbook has no sheet named `Events`."
    )


def test_a_sheet_with_no_heading_row_is_refused() -> None:
    assert _refused(_sheet_bytes()).code == "no_header_row"


def test_a_duplicate_heading_is_refused() -> None:
    refusal = _refused(_sheet_bytes(("id", "ID "), ("P001", "P002")))

    assert refusal.message == (
        "The `Profiles` sheet has two columns named `ID`; please leave one of them and "
        "upload it again."
    )


def test_a_value_in_a_column_with_no_heading_is_refused() -> None:
    refusal = _refused(_sheet_bytes(("id", None), ("P001", "stray")))

    assert refusal.message == (
        "Row 2 of the `Profiles` sheet has a value in a column with no heading; please "
        "give the column a heading or clear it."
    )


def test_too_many_columns_is_refused() -> None:
    headings = tuple(f"c{index}" for index in range(MAX_COLUMN_COUNT + 1))

    assert _refused(_sheet_bytes(headings)).code == "too_many_columns"


def test_an_overlong_cell_is_refused_and_its_heading_is_quoted_safely() -> None:
    refusal = _refused(_sheet_bytes(("na`me\nx",), ("y" * (MAX_CELL_CHARACTERS + 1),)))

    assert refusal.code == "cell_too_long"
    assert "`name x`" in refusal.message


def test_too_many_rows_is_refused_without_reading_them_all() -> None:
    rows = [("id",)] + [(f"P{index}",) for index in range(MAX_DATA_ROW_COUNT + 5)]

    refusal = _refused(_sheet_bytes(*rows))
    assert refusal.message == (
        f"The `Profiles` sheet has more than {MAX_DATA_ROW_COUNT} rows, which is more "
        "than this page reads."
    )


def test_a_formula_is_never_evaluated() -> None:
    raw = _sheet_bytes(("id", "sum"), ("P001", "=1+1"))

    result = _read(raw)
    assert isinstance(result, tuple)
    # ``data_only`` returns Excel's cached value; openpyxl wrote none, so it is blank.
    assert result[0].rows == ((2, {"id": "P001", "sum": ""}),)


# ---------------------------------------------------------------------------
# Bytes and container
# ---------------------------------------------------------------------------


def test_a_file_over_the_byte_cap_is_refused_before_it_is_read() -> None:
    assert _refused(b"PK\x03\x04" + b"\x00" * MAX_UPLOAD_BYTES).code == "file_too_large"


def test_an_empty_upload_is_refused() -> None:
    assert _refused(b"   ").message == "The file is empty."


def test_a_csv_is_refused_with_a_sentence_asking_for_the_workbook() -> None:
    refusal = _refused(b"profile_id,first_name\nP001,Fictional\n")

    assert refusal.code == "not_a_workbook"
    assert ".xlsx" in refusal.message


def test_an_old_binary_xls_is_refused_with_its_own_sentence() -> None:
    assert _refused(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64).code == "old_excel_format"


def test_a_zip_that_is_not_a_workbook_is_refused() -> None:
    assert _refused(_zip({"hello.txt": b"not a workbook"})).code == "unreadable_workbook"


def test_a_truncated_workbook_is_refused() -> None:
    raw = good_workbook()

    assert _refused(raw[: len(raw) // 2]).code == "unreadable_workbook"


# ---------------------------------------------------------------------------
# Zip bomb — checked from the central directory before openpyxl opens anything
# ---------------------------------------------------------------------------


@pytest.fixture
def openpyxl_must_not_open(monkeypatch: pytest.MonkeyPatch) -> None:
    def refuse(*args: object, **kwargs: object) -> None:
        raise AssertionError("openpyxl was asked to open a file the zip guard should refuse")

    monkeypatch.setattr(workbook.openpyxl, "load_workbook", refuse)


@pytest.mark.usefixtures("openpyxl_must_not_open")
def test_a_part_that_inflates_past_the_cap_is_refused_before_openpyxl() -> None:
    bomb = _zip({"xl/worksheets/sheet1.xml": b"\x00" * (MAX_TOTAL_UNCOMPRESSED_BYTES + 1)})

    assert len(bomb) < MAX_UPLOAD_BYTES
    refusal = _refused(bomb)
    assert refusal.code == "workbook_too_large"


@pytest.mark.usefixtures("openpyxl_must_not_open")
def test_many_parts_that_together_inflate_past_the_cap_are_refused() -> None:
    chunk = b"\x00" * MAX_ENTRY_UNCOMPRESSED_BYTES  # each part is within its own cap
    bomb = _zip(
        {
            f"xl/part{index}.xml": chunk
            for index in range(MAX_TOTAL_UNCOMPRESSED_BYTES // MAX_ENTRY_UNCOMPRESSED_BYTES + 1)
        }
    )

    assert _refused(bomb).code == "workbook_too_large"


@pytest.mark.usefixtures("openpyxl_must_not_open")
def test_too_many_parts_is_refused() -> None:
    crowded = _zip({f"xl/part{index}.xml": b"" for index in range(MAX_ZIP_ENTRIES + 1)})

    assert _refused(crowded).code == "workbook_too_large"


@pytest.mark.usefixtures("openpyxl_must_not_open")
def test_an_encrypted_part_is_refused() -> None:
    raw = bytearray(_zip({"xl/workbook.xml": b"<x/>"}))
    # Set the "encrypted" general-purpose flag in the central directory record.
    at = raw.index(b"PK\x01\x02")
    flags = struct.unpack_from("<H", raw, at + 8)[0]
    struct.pack_into("<H", raw, at + 8, flags | 0x1)

    assert _refused(bytes(raw)).code == "workbook_encrypted"


def test_a_part_that_lies_about_its_size_yields_a_refusal_not_more_data() -> None:
    """CPython stops at the declared size, so a lie is a broken file, not a bomb."""
    raw = bytearray(good_workbook())
    name = b"xl/worksheets/sheet2.xml"
    for signature, size_at in ((b"PK\x03\x04", 22), (b"PK\x01\x02", 24)):
        at = raw.index(signature)
        while at != -1:
            name_at = at + (30 if signature == b"PK\x03\x04" else 46)
            if raw[name_at : name_at + len(name)] == name:
                struct.pack_into("<I", raw, at + size_at, 64)
            at = raw.find(signature, at + 4)

    assert _refused(bytes(raw), "Profiles", "Events").code == "unreadable_workbook"


# ---------------------------------------------------------------------------
# XML
# ---------------------------------------------------------------------------


_ENTITY_BOMB = b"""<?xml version="1.0"?>
<!DOCTYPE lolz [<!ENTITY lol "lol"><!ENTITY lol2 "&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;">]>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
<sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>&lol2;</t></is></c></row></sheetData>
</worksheet>"""


def test_an_xml_entity_declaration_is_refused(caplog: pytest.LogCaptureFixture) -> None:
    raw = _replace_part(_sheet_bytes(("id",), ("P001",)), "xl/worksheets/sheet1.xml", _ENTITY_BOMB)

    with caplog.at_level(logging.INFO, logger="smartmatch_domain.exercise.workbook"):
        refusal = _refused(raw)

    assert refusal.code == "unreadable_workbook"
    assert "lol" not in caplog.text


def test_every_upload_is_refused_when_openpyxl_is_not_using_defusedxml(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(workbook.openpyxl, "DEFUSEDXML", False)

    assert _refused(good_workbook(), "Profiles").code == "workbook_reader_unavailable"


def test_openpyxl_is_using_defusedxml_in_this_environment() -> None:
    """The dependency is declared; this proves it is installed and picked up."""
    assert openpyxl.DEFUSEDXML is True


# ---------------------------------------------------------------------------
# Review round 1
# ---------------------------------------------------------------------------


def test_a_stale_sheet_dimension_does_not_drop_rows() -> None:
    """A read-only sheet stops at its declared ``<dimension>``; the reader resets it.

    Found in review: a Profiles sheet declaring ``A1:N55`` over 61 rows parsed as
    54 profiles and was accepted. Rows past a stale declaration must still be read.
    """
    import re

    raw = good_workbook()
    part = "xl/worksheets/sheet2.xml"
    xml = zipfile.ZipFile(io.BytesIO(raw)).read(part)
    stale = re.sub(rb'<dimension ref="[^"]*"\s*/>', b'<dimension ref="A1:N55"/>', xml)
    assert stale != xml, "the fixture's Profiles sheet carries no dimension to falsify"

    result = read_sheets(_replace_part(raw, part, stale), ("Profiles",))

    assert isinstance(result, tuple)
    assert len(result[0].rows) == 60


@pytest.mark.usefixtures("openpyxl_must_not_open")
def test_a_part_compressed_with_another_method_is_refused() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_BZIP2) as archive:
        archive.writestr("xl/workbook.xml", b"<x/>")

    assert _refused(buffer.getvalue()).code == "unreadable_workbook"


def test_two_sheets_that_share_a_name_are_refused_rather_than_guessed() -> None:
    book = openpyxl.Workbook()
    first = book.active
    assert first is not None
    first.title = "Profiles"
    first.append(["id"])
    second = book.create_sheet("profiles ")
    second.append(["id"])
    buffer = io.BytesIO()
    book.save(buffer)

    refusal = _refused(buffer.getvalue())
    assert refusal.code == "duplicate_sheet"
    assert "`Profiles`" in refusal.message


# ---------------------------------------------------------------------------
# Security review of PR #228
# ---------------------------------------------------------------------------

_R_NS = 'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'


def _with_extra_sheets(raw: bytes, count: int, *, shared: bool) -> bytes:
    """``count`` more ``<sheet>`` entries, all on one new part or each on its own."""
    source = zipfile.ZipFile(io.BytesIO(raw))
    parts = {info.filename: source.read(info) for info in source.infolist()}
    body = b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData/></worksheet>'
    sheets, rels = [], []
    for index in range(count):
        rel_id = "rIdExtra" if shared else f"rIdExtra{index}"
        sheets.append(f'<sheet {_R_NS} name="x{index}" sheetId="{index + 50}" r:id="{rel_id}"/>')
        if not shared or index == 0:
            target = "extra.xml" if shared else f"extra{index}.xml"
            rels.append(
                '<Relationship Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                f'relationships/worksheet" Target="/xl/worksheets/{target}" Id="{rel_id}"/>'
            )
            parts[f"xl/worksheets/{target}"] = body
    workbook_xml = parts["xl/workbook.xml"].decode()
    parts["xl/workbook.xml"] = workbook_xml.replace(
        "</sheets>", "".join(sheets) + "</sheets>"
    ).encode()
    rels_xml = parts["xl/_rels/workbook.xml.rels"].decode()
    parts["xl/_rels/workbook.xml.rels"] = rels_xml.replace(
        "</Relationships>", "".join(rels) + "</Relationships>"
    ).encode()
    return _zip(parts)


def test_many_sheets_on_one_part_are_refused_before_any_is_sized() -> None:
    """Finding 1: one part listed many times was parsed once per listing."""
    raw = _with_extra_sheets(_sheet_bytes(("id",), ("P001",)), 2, shared=True)

    assert _refused(raw).code == "too_many_sheets"


def test_more_sheets_than_the_cap_are_refused() -> None:
    raw = _with_extra_sheets(_sheet_bytes(("id",), ("P001",)), MAX_SHEET_COUNT, shared=False)

    assert _refused(raw).code == "too_many_sheets"


def test_a_few_extra_sheets_of_their_own_are_fine() -> None:
    raw = _with_extra_sheets(_sheet_bytes(("id",), ("P001",)), 3, shared=False)

    assert isinstance(_read(raw), tuple)


def test_openpyxl_warnings_are_silenced_so_they_cannot_quote_the_file() -> None:
    """Finding 3: a warning's text quotes the file and is kept in a registry."""
    import warnings

    import openpyxl.reader.workbook as openpyxl_workbook_reader

    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("default")
        # Re-apply the module's own filter inside this block, as it applies at import.
        warnings.filterwarnings("ignore", module=r"openpyxl(\..*)?$")
        warnings.warn_explicit(
            "quoted file content",
            UserWarning,
            "workbook.py",
            1,
            module=openpyxl_workbook_reader.__name__,
        )
    assert caught == []
    source = Path(workbook.__file__).read_text(encoding="utf-8")
    assert 'warnings.filterwarnings("ignore", module=r"openpyxl' in source


def test_every_upload_is_refused_when_openpyxl_would_parse_with_lxml(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Finding 6: with lxml, ``fromstring`` is lxml's, not defusedxml's."""
    monkeypatch.setattr(workbook.openpyxl, "LXML", True)

    assert _refused(good_workbook(), "Profiles").code == "workbook_reader_unavailable"


def test_a_zip_error_of_any_kind_is_one_sentence(monkeypatch: pytest.MonkeyPatch) -> None:
    """Finding 5: ``zipfile`` raises ``NotImplementedError`` for an odd version field."""

    def unsupported(*args: object, **kwargs: object) -> None:
        raise NotImplementedError("zip file version 19.0")

    monkeypatch.setattr(workbook.zipfile, "ZipFile", unsupported)

    assert _refused(good_workbook()).code == "unreadable_workbook"


def test_a_sheet_title_is_never_quoted_from_the_file() -> None:
    """Finding 4: the sentence names the layout's sheet, not the workbook's title."""
    refusal = _refused(_sheet_bytes(title="profiles`\n"), "Profiles")

    assert refusal.message == "The `Profiles` sheet has no heading row naming its columns."
